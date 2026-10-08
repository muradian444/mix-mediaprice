# -*- coding: utf-8 -*-
"""Պայմանագրերի (0046 / 0055 / 0056) լրացում Word շաբլոններից՝ ձևաչափը պահպանելով."""
import copy
import re
import shutil
import subprocess
import tempfile
import threading
from datetime import date
from pathlib import Path

import docx

from numwords import money, words

BASE = Path(__file__).resolve().parent.parent   # նախագծի արմատը (app/-ից մեկ մակարդակ վեր)
TPL = BASE / "templates"

MONTHS_GEN = ["հունվարի", "փետրվարի", "մարտի", "ապրիլի", "մայիսի", "հունիսի", "հուլիսի",
              "օգոստոսի", "սեպտեմբերի", "հոկտեմբերի", "նոյեմբերի", "դեկտեմբերի"]
MONTHS_NOM = ["հունվար", "փետրվար", "մարտ", "ապրիլ", "մայիս", "հունիս", "հուլիս",
              "օգոստոս", "սեպտեմբեր", "հոկտեմբեր", "նոյեմբեր", "դեկտեմբեր"]


# ---------------------------------------------------------------- helpers
def sub_par(p, pattern, repl, flags=0):
    """Regex-փոխարինում պարբերության մեջ՝ run-երի սահմաններից անկախ (ձևաչափը պահվում է)."""
    runs = p.runs
    full = "".join(r.text for r in runs)
    m = re.search(pattern, full, flags)
    if not m:
        return False
    new = m.expand(repl) if isinstance(repl, str) and "\\" in repl else repl
    start, end = m.span()
    pos = 0
    placed = False
    for r in runs:
        t = r.text
        a, b = pos, pos + len(t)
        pos = b
        if b <= start or a >= end:
            if not (a == start == end):
                continue
        head = t[:max(0, start - a)]
        tail = t[max(0, end - a):] if end < b else ""
        if not placed and a <= start < b or (not placed and start == b and False):
            r.text = head + new + tail
            placed = True
        elif a < end:
            r.text = tail
    if not placed and runs:  # empty match edge case
        runs[-1].text += new
    return True


def all_pars(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield p


def sub_all(doc, pattern, repl, flags=0):
    n = 0
    for p in all_pars(doc):
        if sub_par(p, pattern, repl, flags):
            n += 1
    return n


def parse_date(s):
    s = s.strip()
    if s.upper() == "TODAY":
        return date.today()
    d, m, y = re.split(r"[.\-/]", s)
    return date(int(y), int(m), int(d))


def add_year(d):
    try:
        return d.replace(year=d.year + 1)
    except ValueError:
        return d.replace(year=d.year + 1, day=28)


def fmt(d):
    return d.strftime("%d.%m.%Y")


def gen(name):
    """Տնօրենի անունը սեռական հոլովով (պարզ կանոն՝ + ի)."""
    name = name.strip()
    return name if name.endswith("ի") else name + "ի"


def cap(s):
    return s[:1].upper() + s[1:]


# ---------------------------------------------------------------- dates
def fill_dates(doc, d, lower_month=False):
    dd, mm, yy = f"{d.day:02d}", d.month - 1, d.year
    gen_m = MONTHS_GEN[mm]
    nom_m = MONTHS_NOM[mm]
    sub_all(doc, r"«\s*\d{1,2}\s*»\s*\S+\s+\d{4}",
            f"«{dd}» {gen_m if lower_month else cap(gen_m)} {yy}")
    sub_all(doc, r"\d{4}\s*ԹՎԱԿԱՆԻ\s*\S+\s*\d{1,2}-ԻՆ",
            f"{yy} ԹՎԱԿԱՆԻ {gen_m.upper()} {dd}-ԻՆ")
    sub_all(doc, r"\d{4}\s*թվականի\s*\S+\s*\d{1,2}-ին",
            f"{yy} թվականի {gen_m} {dd}-ին")


def fill_nominative_header_0056(doc, d):
    sub_all(doc, r"«\s*\d{1,2}\s*»\s*Սեպտեմբեր\s+\d{4}",
            f"«{d.day:02d}» {cap(MONTHS_NOM[d.month - 1])} {d.year}")


# ---------------------------------------------------------------- parties
def fill_intro(doc, data):
    comp, director = data["company"], gen(data["director"])
    for p in doc.paragraphs:
        txt = p.text
        if "ԱՎԵԼԱՑՆԵԼ" in txt or "Ավելացնել" in txt:
            sub_par(p, r"ԱՎԵԼԱՑՆԵԼ\s*-\s*ան\s*ի?", director) or \
                sub_par(p, r"Ավելացնել\s*-\s*ան\s*ի", director)
            sub_par(p, r"ԱՎԵԼԱՑՆԵԼ", comp)
            break


def bank_line(data):
    """«Բանկ, հաշվեհամար»՝ առանց ավելորդ ստորակետի, եթե դաշտերից մեկը (կամ երկուսն էլ) դատարկ է."""
    return ", ".join(x for x in (data.get("bank", ""), data.get("account", "")) if x)


def _client_par(p, data, is_first_tin_par=None):
    t = p.text.strip()
    if not t:
        return
    if t.startswith("«") and "ՍՊԸ" in t:
        sub_par(p, r"«[^»]*»", f"«{data['company']}»")
    elif t.startswith("ՀՎՀՀ"):
        sub_par(p, r"՝[ \t]*[^\n]*", f"՝ {data['tin']}")
    elif t.startswith("ՀՀ"):
        sub_par(p, r"՝[ \t]*[^\n]*", f"՝ {bank_line(data)}")
    elif t.startswith("Բանկային հաշվեհամար"):
        sub_par(p, r"՝[ \t]*[^\n]*", f"՝ {bank_line(data)}")
    else:
        if "Հասցե" in t:
            sub_par(p, r"Հասցե\s*՝[ \t]*[^\n]*", f"Հասցե՝ {data['legal_address']}")
        if "Էլ" in t:
            sub_par(p, r"(Էլ[^\n՝`]*[՝`])[ \t]*[^\n]*", r"\1 " + data.get("email", "").replace("\\", ""))
        if t.startswith("Տնօրեն"):
            sub_par(p, r"Տնօրեն\s*՝[ \t]*[^\n_]*", f"Տնօրեն՝ {data['director']}")


def fill_party_tables(doc, data, add_bank_row=False):
    for t in doc.tables:
        if not t.rows or not t.rows[0].cells or "ԿԱՏԱՐՈՂ" not in t.rows[0].cells[0].text:
            continue
        for row in t.rows[1:]:
            if len(row.cells) < 2:
                continue
            cell = row.cells[1]
            for p in list(cell.paragraphs):
                first = p.text.strip().startswith("ՀՎՀՀ")
                _client_par(p, data)
                if add_bank_row and first and bank_line(data):
                    newp = copy.deepcopy(p._p)
                    p._p.addnext(newp)
                    from docx.text.paragraph import Paragraph
                    np_ = Paragraph(newp, p._parent)
                    sub_par(np_, r"ՀՎՀՀ\s*՝?[ \t]*", "ՀՀ՝ ")
                    sub_par(np_, r"՝[ \t]*[^\n]*", f"՝ {bank_line(data)}")


# ---------------------------------------------------------------- addresses
def fill_address_list(doc, anchor_indices, addresses):
    """anchor_indices – [Հասցե 1], [Հասցե 2] պարբերությունների ինդեքսները (շաբլոնում)."""
    pars = doc.paragraphs
    first = pars[anchor_indices[0]]
    anchors = [pars[i] for i in anchor_indices]
    base = first._p
    # կլոնավորում ենք առաջին պարբերությունը անհրաժեշտ քանակով
    prev = anchors[-1]._p
    new_elems = []
    for _ in addresses:
        el = copy.deepcopy(base)
        prev.addnext(el)
        prev = el
        new_elems.append(el)
    from docx.text.paragraph import Paragraph
    for el, addr in zip(new_elems, addresses):
        para = Paragraph(el, first._parent)
        runs = para.runs
        if runs:
            runs[0].text = f"{addresses.index(addr) + 1}. {addr}" if False else ""
    for i, (el, addr) in enumerate(zip(new_elems, addresses), 1):
        para = Paragraph(el, first._parent)
        runs = para.runs
        for r in runs[1:]:
            r.text = ""
        if runs:
            runs[0].text = f"{i}. {addr}"
    for a in anchors:
        a._p.getparent().remove(a._p)


# ---------------------------------------------------------------- templates
def build_0055(d):
    doc = docx.Document(TPL / "0055.docx")
    cd = parse_date(d["date"])
    fill_dates(doc, cd)
    fill_intro(doc, d)
    fill_party_tables(doc, d)
    sub_all(doc, r"\[Խանութի անվանումը\]", d["shop"])
    sub_all(doc, r"3 կամ 4", d["points"])
    n, price = int(d["addr_count"]), int(d["price"])
    total = n * price
    sub_all(doc, r"8000 \(ութ հազար\)", f"{money(price)} ({words(price)})")
    sub_all(doc, r"17 հասցեի համար", f"{n} հասցեի համար")
    sub_all(doc, r"136,000 դրամ \(Հարյուր երեսունվեց հազար դրամ\)",
            f"{money(total)} դրամ ({cap(words(total))} դրամ)")
    s, l = int(d["upto120"]), int(d["over120"])
    sub_all(doc, r"6,000 \(վեց հազար\)", f"{money(s)} ({words(s)})")
    sub_all(doc, r"12,000 \(տասներկու հազար\)", f"{money(l)} ({words(l)})")
    months = int(d["months"])
    sub_all(doc, r"3 \(երեք\) ամիսը մեկ", f"{months} ({words(months)}) ամիսը մեկ")
    if months != 3:
        sub_all(doc, r"եռամսյակի", "ժամանակահատվածի")
    sub_par_text = doc.paragraphs
    # հասցեներ
    idx = [i for i, p in enumerate(doc.paragraphs) if re.match(r"\s*\[\s*Հասցե\s*\d\]", p.text)]
    sub_all(doc, r"\(Վարչական շրջան և հասցե\)՝\s*\n?Մարզեր\s*\(Վարչական շրջան, քաղաք և հասցե\)՝", "՝")
    fill_address_list(doc, idx, d["addresses"])
    return doc


def build_0046(d):
    doc = docx.Document(TPL / "0046.docx")
    cd = parse_date(d["date"])
    fill_dates(doc, cd)
    fill_intro(doc, d)
    fill_party_tables(doc, d, add_bank_row=True)
    sub_all(doc, r"\[Խանութի անվանումը\]", d["shop"])
    s, l = int(d["upto120"]), int(d["over120"])
    sub_all(doc, r"6,000 \(վեց հազար\)", f"{money(s)} ({words(s)})")
    sub_all(doc, r"12,000 \(տասներկու հազար\)", f"{money(l)} ({words(l)})")
    months = int(d["months"])
    sub_all(doc, r"3 \(երեք\) ամիսը մեկ", f"{months} ({words(months)}) ամիսը մեկ")
    if months != 3:
        sub_all(doc, r"եռամսյակի", "ժամանակահատվածի")
    idx = [i for i, p in enumerate(doc.paragraphs) if re.match(r"\s*\[\s*Հասցե\s*\d\]", p.text)]
    fill_address_list(doc, idx, d["addresses"])
    return doc


def build_0056(d):
    doc = docx.Document(TPL / "0056.docx")
    cd = parse_date(d["date"])
    end = add_year(cd)
    fill_dates(doc, cd, lower_month=True)
    fill_nominative_header_0056(doc, cd)
    fill_intro(doc, d)
    fill_party_tables(doc, d)
    price = int(d["price"])
    sub_all(doc, r"70,000 \(յոթանասուն հազար\)", f"{money(price)} ({words(price)})")
    sub_all(doc, r"70,000 ՀՀ դրամ", f"{money(price)} ՀՀ դրամ")
    sub_all(doc, r"840,000 ՀՀ դրամ", f"{money(price * 12)} ՀՀ դրամ")
    sub_all(doc, r"\d{2}\.\d{2}\.\d{4}\s*–\s*\d{2}\.\d{2}\.\d{4}", f"{fmt(cd)} – {fmt(end)}")
    sub_all(doc, r"(?<=սահմանվում՝)[ \t]*[^\n]*", f" {d['location'].rstrip('։.')}։")
    sub_all(doc, r"24 հեռարձակում", f"{int(d['spots'])} հեռարձակում")
    sub_all(doc, r"30վրկ", f"{int(d['duration'])}վրկ")
    if d.get("remove44"):
        pars = doc.paragraphs
        for i, p in enumerate(pars):
            if p.text.strip().startswith("4․4․") or p.text.strip().startswith("4.4."):
                p._p.getparent().remove(p._p)
                break
        for p in doc.paragraphs:
            t = p.text.strip()
            if t.startswith("4.5."):
                sub_par(p, r"4\.5\.", "4.4.")
            elif t.startswith("4.6."):
                sub_par(p, r"4\.6\.", "4.5.")
    return doc


BUILDERS = {"0055": build_0055, "0046": build_0046, "0056": build_0056}


def make_contract(kind, data, out_dir):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = BUILDERS[kind](data)
    safe = re.sub(r"[^\w\-]+", "_", data["company"], flags=re.U)[:40]
    path = out_dir / f"Payman_{kind}_{safe}_{parse_date(data['date']).strftime('%d%m%Y')}.docx"
    doc.save(path)
    return path


_PDF_LOCK = threading.Lock()     # Word/LibreOffice-ը միաժամանակ երկու փոխարկում լավ չի անում


def _soffice():
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        for c in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                  r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
            if Path(c).exists():
                return c
    return exe


def _word_pdf(docx_path, out):
    """Microsoft Word (COM). Սերվերի աշխատանքային թելում պետք է CoInitialize, այլապես միշտ ձախողվում էր."""
    co = None
    try:
        import pythoncom
        pythoncom.CoInitialize()
        co = pythoncom
    except Exception:  # noqa — pywin32 չկա՝ docx2pdf-ը ինքը կփորձի
        pass
    try:
        import contextlib
        import io
        from docx2pdf import convert
        with contextlib.redirect_stderr(io.StringIO()):   # tqdm-ի առաջընթացը չխառնի ադմինի կոնսոլը
            convert(str(docx_path), str(out))
        return out if out.exists() else None
    except Exception:  # noqa
        return None
    finally:
        if co:
            try:
                co.CoUninitialize()
            except Exception:  # noqa
                pass


def to_pdf(docx_path):
    """docx -> pdf (LibreOffice կամ Microsoft Word). Վերադարձնում է pdf-ի ճանապարհը կամ None."""
    exe = _soffice()
    out = Path(docx_path).with_suffix(".pdf")
    with _PDF_LOCK:
        if not exe:
            return _word_pdf(docx_path, out)
        try:
            subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(out.parent), str(docx_path)],
                           check=True, capture_output=True, timeout=300)
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError):
            return None   # PDF չստացվեց՝ Word տարբերակը մնում է, հավելվածը չի ընկնում
    return out if out.exists() else None


_PDF_AVAIL = []


def pdf_available():
    """-> (True/False, բացատրություն)՝ կարո՞ղ ենք DOCX-ը PDF դարձնել այս համակարգչում (մեկ անգամ ստուգվում է)."""
    if _PDF_AVAIL:
        return _PDF_AVAIL[0]
    res = (False, "LibreOffice կամ Microsoft Word տեղադրված չէ")
    if _soffice():
        res = (True, "LibreOffice")
    else:
        try:
            import docx2pdf  # noqa: F401
            import winreg
            winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "Word.Application")
            res = (True, "Microsoft Word")
        except Exception:  # noqa
            pass
    _PDF_AVAIL.append(res)
    return res
