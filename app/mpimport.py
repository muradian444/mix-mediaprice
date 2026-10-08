# -*- coding: utf-8 -*-
"""Պատրաստի մեդիա պլանի ֆայլի ընթերցում (PDF / Excel / CSV / Word) -> պատվիրատու, ժամանակահատված,
եթերացանկ (ժամեր) և հասցեներ: ԱԿՏ-ում հասցեներով անջատված օրերը հաշվելու համար."""
import re
from datetime import date
from pathlib import Path

import act

SUPPORTED = (".pdf", ".xlsx", ".xlsm", ".csv", ".txt", ".docx")
COMPANY_MARK = "«Միքս Մեդիա»"
PERIOD_RE = re.compile(r"(?<!\d)(\d{1,2})[./․](\d{1,2})[./․](\d{4}|\d{2})(?!\d)\s*[-–—]\s*"
                       r"(\d{1,2})[./․](\d{1,2})[./․](\d{4}|\d{2})(?!\d)")     # 24.06.2026 - 23.07.2026 / 24.06.26 - 23.07.26
ADDR_HEAD = ("հասցե", "адрес", "address", "shop", "store", "խանութ", "магазин")
NET_HEAD = ("ցանց", "сеть", "network", "net")
CLIENT_LABELS = ("պատվիրատու", "հաճախորդ", "клиент", "заказчик", "client", "customer")


def _clean(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def _period(text):
    m = PERIOD_RE.search(text or "")
    if not m:
        return None, None
    y1, y2 = (int(y) + (2000 if len(y) == 2 else 0) for y in (m[3], m[6]))
    try:
        a, b = date(y1, int(m[2]), int(m[1])), date(y2, int(m[5]), int(m[4]))
    except ValueError:
        return None, None
    return (a.isoformat(), b.isoformat()) if a <= b else (None, None)


def _slot(s):
    """'9:20' / '09:20:00' / datetime.time -> '9:20' (այլապես None)."""
    if hasattr(s, "hour") and hasattr(s, "minute"):
        return f"{s.hour}:{s.minute:02d}"
    m = re.fullmatch(r"\s*(\d{1,2}):(\d{2})(?::\d{2})?\s*", str(s or ""))
    if m and int(m[1]) < 24 and int(m[2]) < 60:
        return f"{int(m[1])}:{m[2]}"
    return None


def _result(src, client="", start=None, end=None, slots=None, addresses=None, warnings=None, schedules=None):
    seen, addrs = set(), []
    for a in addresses or []:
        addr = _clean(a.get("addr"))
        if addr and addr.lower() not in seen:
            seen.add(addr.lower())
            row = dict(net=_clean(a.get("net")), addr=addr)
            if a.get("slots"):
                row["slots"] = list(a["slots"])
            addrs.append(row)
    slots = list(dict.fromkeys(slots or []))
    return dict(source=src, client=_clean(client), start=start or "", end=end or "", slots=slots,
                addresses=addrs, warnings=list(warnings or []), schedules=schedules or [])


# ------------------------------------------------------------------ Mix Media-ի ստանդարտ մեդիա պլան (Excel / PDF)
# Մի քանի բաժին՝ «<Ցանց> 57 հասցե» + եթերացանկ («9:00 1 1 1 … 29» տողեր, ժամը կարող է լինել «9։20»),
# վերջում՝ «Աուդիոհոլովակների հեռարձակման հասցեները»՝ ցանցի անուն + համարակալված հասցեներ:
# Ամեն ցանց ստանում է ԻՐ բաժնի եթերացանկը (օր.՝ Երևան Սիթի 9:00–23:30, Մեգա Մոլ 9:30–21:35/25):
TIME_ROW = re.compile(r"^(\d{1,2})\s*[:։]\s*(\d{2})(?:[:։]\d{2})?(?=\s|$)")
SEC_NET = re.compile(r"^(.+?)\s+(\d{1,4})\s+հասցե$")
NUM_ADDR = re.compile(r"^(\d{1,3})\s*[.․)]?\s*([^\d\s.․)].*|\d+\s*[^\d\s].*)$")
NETLIKE = re.compile(r"ցանց|սուպերմարկետ|մարկետ|խանութ|^մոլեր$|market|mall|сеть|магазин", re.I)
SKIP_ADDR = ("պատվիրատու", "հեռարձակման ամս", "աուդիոհոլովակների", "ընդհանուր հեռարձակում", "հասցեների քանակ")
CLIENT_RE = re.compile(r"^(?:պատվիրատու|հաճախորդ|заказчик|клиент|client|customer)\s*[`՝:'’]?\s*(.+)$", re.I)


def _is_grid(lines):
    """Եթերացանկի ամբողջ տողը մեկ տողում է («9:00 1 1 1 … 29»)՝ ստանդարտ (Excel-ից) մեդիա պլան."""
    return sum(1 for s in lines if TIME_ROW.match(s) and len(s.split()) >= 4) >= 3


def _client_of(lines):
    for s in lines[:80]:
        m = CLIENT_RE.match(_clean(s))
        if m and len(m[1]) > 1:
            c = re.sub(r'"([^"]+)"', r"«\1»", m[1].strip(' `՝:'))
            return re.sub(r"^ՍՊԸ\s*«", "«", c)
    return ""


def _words(name):
    import monitor
    return monitor.net_words(name)


def _net_match(a, b):
    """Ցանցերի անունները նույնն են (առաջին բառով)՝ «Էտալոն» ~ «Էտալոն սուպերմարկետ»."""
    wa, wb = _words(a), _words(b)
    if not wa or not wb:
        return False
    x, y = wa[0], wb[0]
    return x == y or (min(len(x), len(y)) >= 4 and x[:4] == y[:4])


def parse_sections(lines):
    """-> (sections[{nets:[(անուն, քանակ)], slots}], groups[(ցանց, [հասցեներ])])."""
    sections, cur, in_addr = [], None, False
    groups, g, n = [], None, 0
    for raw in lines:
        s = _clean(raw)
        if not s or _is_footer(s):
            continue
        low = s.lower()
        if "հեռարձակման հասցեները" in low or "адреса вещания" in low:
            in_addr, g, n = True, None, 0
            continue
        if not in_addr:
            m = TIME_ROW.match(s)
            if m and int(m[1]) < 24 and int(m[2]) < 60:
                if cur is None:
                    cur = dict(nets=[], slots=[])
                    sections.append(cur)
                cur["slots"].append(f"{int(m[1])}:{m[2]}")
                continue
            m = SEC_NET.match(s)
            if m:
                if cur is None or cur["slots"]:          # նոր բաժին (նախորդի եթերացանկից հետո)
                    cur = dict(nets=[], slots=[])
                    sections.append(cur)
                cur["nets"].append((m[1].strip(" ,:՝"), int(m[2])))
            continue
        if any(low.startswith(k) for k in SKIP_ADDR) or s in ("N", "Հասցե", "№"):
            continue
        m = NUM_ADDR.match(s)
        if m and (int(m[1]) == n + 1 or int(m[1]) == 1):
            if g is None or (int(m[1]) == 1 and n > 0 and g[1]):   # համարակալումը նորից 1-ից՝ նոր խումբ
                g = ["", []]
                groups.append(g)
            n = int(m[1])
            g[1].append(m[2].strip(" ։.,"))
            continue
        if NETLIKE.search(s) or g is None or not g[1]:     # ցանցի վերնագիր
            g, n = [s.strip(" :՝"), []], 0
            groups.append(g)
        else:                                               # նախորդ հասցեի շարունակություն
            g[1][-1] = (g[1][-1] + " " + s).strip()
    return [x for x in sections if x["slots"]], [x for x in groups if x[1]]


def _section_for(net, addrs, sections):
    for sec in sections:
        if any(_net_match(nm, net) for nm, _ in sec["nets"]):
            return sec
    flat_addrs = [_flat(a) for a in addrs]
    for sec in sections:          # «Մեգա Մոլ 1 հասցե» բաժին ↔ «Մոլեր» ցանցի «Մեգա մոլ» հասցե
        if any(len(_flat(nm)) >= 4 and any(_flat(nm) in fa for fa in flat_addrs) for nm, _ in sec["nets"]):
            return sec
    same = [sec for sec in sections if any(c == len(addrs) for _, c in sec["nets"])]
    if len(same) == 1:
        return same[0]
    return sections[0] if len(sections) == 1 else None


def from_sections(lines, src):
    """Ստանդարտ մեդիա պլան -> _result (յուրաքանչյուր հասցե՝ իր բաժնի եթերացանկով) կամ None."""
    if not _is_grid(lines):
        return None
    sections, groups = parse_sections(lines)
    if not sections or not groups:
        return None
    start, end = _period("\n".join(_clean(x) for x in lines))
    addrs, warn = [], []
    for net, lst in groups:
        sec = _section_for(net, lst, sections)
        if sec is None:
            warn.append(f"«{net or '?'}»՝ եթերացանկը չգտնվեց, կօգտագործվի ընդհանուրը")
        for a in lst:
            addrs.append(dict(net=net, addr=a, slots=sec["slots"] if sec else None))
    main = max(sections, key=lambda x: sum(c for _, c in x["nets"]) or len(x["slots"]))
    schedules = [dict(nets=[nm for nm, _ in sec["nets"]], count=sum(c for _, c in sec["nets"]), slots=sec["slots"])
                 for sec in sections]
    return _result(src, _client_of(lines), start, end, main["slots"], addrs, warn, schedules)


# ------------------------------------------------------------------ PDF
def _is_footer(s):
    return s.startswith(COMPANY_MARK) and "Երևան" in s and ("Լևոնյան" in s or "Հեռ" in s)


def _pdf_lines(path):
    """Էջերի տողերը՝ առանց էջատակի («Միքս Մեդիա» ՍՊԸ, … Հեռ.) և դրան հաջորդող էջի համարի
    (Word-ից փոխարկված PDF-ում էջատակը կարող է ընկնել էջի մեջտեղում)."""
    from pypdf import PdfReader
    pages = []
    for p in PdfReader(str(path)).pages:
        raw = [_clean(x) for x in (p.extract_text() or "").splitlines()]
        raw = [x for x in raw if x]
        lines, skip_num = [], False
        for x in raw:
            if _is_footer(x):
                skip_num = True
                continue
            if skip_num and x.isdigit():
                skip_num = False
                continue
            skip_num = False
            lines.append(x)
        pages.append(lines)
    return pages


def _known_addresses():
    try:
        import store
        return [(n["name"], a) for n in store.networks() for a in n["addresses"]]
    except Exception:  # noqa
        return []


def parse_pdf(path):
    pages = _pdf_lines(path)
    text_lines = [x for p in pages for x in p]
    std = from_sections(text_lines, "pdf")
    if std:
        return std
    client, start, end = "", None, None
    for i, s in enumerate(text_lines):
        if s == "ՕՐԵՐ":   # վերնագրի շերտ՝ ՊԱՏՎԻՐԱՏՈՒ | ԺԱՄԱՆԱԿԱՀԱՏՎԱԾ | ՕՐԵՐ
            buf = []
            for x in text_lines[i + 1:i + 6]:
                if PERIOD_RE.search(x):
                    start, end = _period(x)
                    break
                buf.append(x)
            client = " ".join(buf)
            break
    if not start:
        start, end = _period("\n".join(text_lines))
    slots = []
    for s in text_lines:
        t = _slot(s)
        if t:
            slots.append(t)
    # հասցեների էջ(եր)
    addrs, started = [], False
    rows = []
    multi = False
    for s in text_lines:
        if s.startswith("Ընդհանուր հեռարձակումներ"):
            break
        if not started:
            if s == "Հասցե" or s.startswith("Հասցե "):
                started = True
            continue
        if s in ("N", "Հասցե") or s.startswith("Հոլովակ"):
            multi = multi or s.startswith("Հոլովակ")
            continue
        rows.append(s)
    if started:
        addrs = _split_rows(rows, multi)
    if not addrs:   # այլ PDF՝ փնտրում ենք համակարգում հայտնի հասցեները
        flat = _clean(" ".join(text_lines)).lower()
        addrs = [dict(net=n, addr=a) for n, a in _known_addresses() if _clean(a).lower() in flat]
    warn = []
    if not addrs:
        warn.append("Հասցեներ չգտա ֆայլում")
    return _result("pdf", client, start, end, slots, addrs, warn)


def _split_rows(lines, multi_clip):
    """[N, հասցե..., (հոլովակներ), (ցանցի անուն), N, ...] -> [{net, addr}]."""
    out, net, i, n = [], "", 0, 0
    clip_re = re.compile(r"^\d+(\s*,\s*\d+)*$")
    while i < len(lines):
        s = lines[i]
        if s.isdigit() and (int(s) == n + 1 or s == "1"):
            n = int(s)
            i += 1
            buf = []
            while i < len(lines):
                nxt = lines[i]
                if nxt.isdigit() and (int(nxt) == n + 1 or nxt == "1"):
                    break
                buf.append(nxt)
                i += 1
            more = bool(i < len(lines) and lines[i] == "1" and n >= 1)   # հաջորդ ցանցի սկիզբ
            nxt_net = ""
            if more and len(buf) >= 2 and not clip_re.match(buf[-1]):   # [հասցե, (հոլովակներ), հաջորդ ցանց]
                nxt_net = buf.pop()
            if buf and len(buf) >= 2 and clip_re.match(buf[-1]) and (multi_clip or "," in buf[-1]):
                buf.pop()                                                 # «1, 2, 3, 4»՝ հոլովակների սյունակ
            out.append(dict(net=net, addr=" ".join(buf)))
            net = nxt_net or net
        else:
            net = s    # առաջին ցանցի վերնագիրը
            i += 1
    return out


# ------------------------------------------------------------------ Excel / CSV
def _cell_text(c):
    """Excel-ի բջիջ -> տեքստ (ժամը՝ «9:20», ամսաթիվը՝ «03.10.2026», 1.0՝ «1»)."""
    from datetime import datetime, time
    if c is None:
        return ""
    if isinstance(c, datetime):
        if c.year <= 1900 and (c.hour or c.minute):
            return f"{c.hour}:{c.minute:02d}"
        return f"{c:%d.%m.%Y}"
    if isinstance(c, date):
        return f"{c:%d.%m.%Y}"
    if isinstance(c, time):
        return f"{c.hour}:{c.minute:02d}"
    if isinstance(c, float):
        if c.is_integer():
            return str(int(c))
        if 0 < c < 1:                      # Excel-ի ժամ՝ օրվա մաս
            m = int(round(c * 1440))
            return f"{m // 60}:{m % 60:02d}"
    return _clean(c)


def _row_lines(rows):
    return [s for s in (" ".join(x for x in (_cell_text(c) for c in r) if x) for r in rows) if s]


def parse_table(path):
    rows = act.read_rows(path)
    std = from_sections(_row_lines(rows), "table")
    if std:
        return std
    head_i = addr_c = net_c = None
    for i, r in enumerate(rows[:60]):
        for j, c in enumerate(r):
            if any(k in str(c or "").lower() for k in ADDR_HEAD) and len(_clean(c)) <= 30:
                head_i, addr_c = i, j
                break
        if head_i is not None:
            for j, c in enumerate(rows[head_i]):
                if any(k in str(c or "").lower() for k in NET_HEAD) and len(_clean(c)) <= 30:
                    net_c = j
            break
    top = rows[:(head_i if head_i is not None else 12)]
    start, end = _period(" ".join(_clean(c) for r in top for c in r if c))
    client = ""
    for r in rows[:25]:
        cells = [_clean(c) for c in r]
        for j, c in enumerate(cells):
            low = c.lower().rstrip(":՝ ")
            if low in CLIENT_LABELS or any(low.startswith(k) for k in CLIENT_LABELS) and len(c) < 25:
                client = next((x for x in cells[j + 1:] if x), "")
                if not client and ":" in c:
                    client = c.split(":", 1)[1].strip()
                break
        if client:
            break
    slots = []
    for r in rows[:(head_i if head_i is not None else 60)] + ([rows[head_i]] if head_i is not None else []):
        for c in r:
            t = _slot(c)
            if t:
                slots.append(t)
    addrs = []
    if head_i is not None:
        cur_net = ""
        for r in rows[head_i + 1:]:
            if addr_c >= len(r) or not _clean(r[addr_c]):
                continue
            if net_c is not None and net_c < len(r) and _clean(r[net_c]):
                cur_net = _clean(r[net_c])
            addrs.append(dict(net=cur_net, addr=r[addr_c]))
    warn = [] if addrs else ["Հասցեների սյունակ («Հասցե» / «Адрес») չգտա"]
    return _result("table", client, start, end, slots, addrs, warn)


# ------------------------------------------------------------------ Word
def parse_docx(path):
    import docx
    d = docx.Document(str(path))
    std = from_sections([p.text for p in d.paragraphs]
                        + _row_lines([[c.text for c in r.cells] for t in d.tables for r in t.rows]), "docx")
    if std:
        return std
    text = "\n".join(p.text for p in d.paragraphs)
    start, end = _period(text)
    client = ""
    for p in d.paragraphs:
        low = p.text.strip().lower()
        if any(low.startswith(k) for k in CLIENT_LABELS) and re.search(r"[՝:]", p.text):
            client = re.split(r"[՝:]", p.text, 1)[1].strip()
            break
    addrs, slots = [], []
    for t in d.tables:
        if not t.rows:
            continue
        heads = [_clean(c.text).lower() for c in t.rows[0].cells]
        col = next((j for j, h in enumerate(heads) if any(k in h for k in ADDR_HEAD)), None)
        if col is None:
            continue
        for r in t.rows[1:]:
            if col < len(r.cells) and _clean(r.cells[col].text):
                addrs.append(dict(net="", addr=r.cells[col].text))
    if not addrs:
        addrs = [dict(net=n, addr=a) for n, a in _known_addresses() if _clean(a).lower() in _clean(text).lower()]
    for line in text.splitlines():
        t = _slot(line)
        if t:
            slots.append(t)
    return _result("docx", client, start, end, slots, addrs, [] if addrs else ["Հասցեներ չգտա ֆայլում"])


def _flat(s):
    return re.sub(r"[\s.,;:·․\-–—«»\"'()/\\]+", "", str(s or "").lower())


def snap(addresses):
    """Ֆայլից կարդացած հասցեները -> Կարգավորումների ցանցերի ճշգրիտ հասցեները (ցանց + հասցե):
    Նախ՝ ճիշտ համընկնում, հետո՝ տան համարով/փողոցով/հարկով (ամեն հայտնի հասցե՝ միայն մեկ անգամ).
    Չգտնվածները մնում են ինչպես կան. -> (addresses, չգտնված քանակ)."""
    known = _known_addresses()
    if not known:
        return list(addresses or []), 0
    import monitor
    by_flat = {}
    for n, a in known:
        by_flat.setdefault(_flat(a), []).append((n, a))
    res, used, rest = {}, set(), []
    for i, x in enumerate(addresses or []):
        cands = [k for k in by_flat.get(_flat(x["addr"]), []) if k not in used]
        hit = next((k for k in cands if x.get("net") and _net_match(k[0], x["net"])), cands[0] if cands else None)
        if hit:
            res[i] = hit
            used.add(hit)
        else:
            rest.append(i)
    if rest:
        keys = [k for k in known if k not in used]
        targets = [dict(net=addresses[i].get("net", ""), addr=addresses[i]["addr"]) for i in rest]
        match = monitor.assign(targets, keys, monitor._all_net_words())
        for i in rest:
            hit = match.get((addresses[i].get("net", ""), addresses[i]["addr"]))
            if hit and hit not in used:
                res[i] = hit
                used.add(hit)
    out, unknown = [], 0
    for i, x in enumerate(addresses or []):
        row = dict(net=res[i][0], addr=res[i][1]) if i in res else dict(net=x.get("net", ""), addr=x["addr"])
        if x.get("slots"):
            row["slots"] = x["slots"]
        unknown += i not in res
        out.append(row)
    return out, unknown


def parse(path):
    """-> {source, client, start, end, slots, addresses[{net, addr}], warnings}՝ հասցեները ճշտված ցանցերով."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        res = parse_pdf(path)
    elif ext == ".docx":
        res = parse_docx(path)
    elif ext in (".xlsx", ".xlsm", ".csv", ".txt"):
        res = parse_table(path)
    else:
        raise ValueError(f"«{ext}» ձևաչափը չի աջակցվում: Ուղարկեք PDF, Excel, CSV կամ Word")
    if res["addresses"]:
        res["addresses"], unknown = snap(res["addresses"])
        if unknown:
            res["warnings"].append(f"Ցանցերում չգտնվեց {unknown} հասցե (թողնված է ինչպես ֆայլում)")
        if not any(a.get("slots") for a in res["addresses"]):
            res["schedules"] = []
    return res
