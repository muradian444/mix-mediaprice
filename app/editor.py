# -*- coding: utf-8 -*-
"""✏️ Խմբագրիչ / Редактор: պատրաստի ֆայլը վերբեռնում ես և փոխում հենց կայքում.

  • DOCX — պարբերություններն ու աղյուսակների վանդակները խմբագրվում են, ձևաչափը (տառատեսակ, թավ, …) մնում է:
  • XLSX / CSV — աղյուսակ՝ վանդակները փոխվում են, բանաձևերը պահպանվում են:
  • TXT — տեքստ:
  • PDF / նկար — տեքստը հանվում է և դառնում խմբագրելի Word (սկանի համար՝ Claude-ով):
  • Մեր հավելվածի ստեղծած պայմանագիր / մեդիա պլան / ԱԿՏ / ԿՊ — կարելի է բացել ՁԵՎՈՒՄ (docmeta)
    և ստեղծել նորից փոխված տվյալներով:"""
import copy
import csv
import difflib
import io
import json
import re
import shutil
from datetime import date, datetime
from pathlib import Path

import docmeta
from config import OUT, UPL
from store import new_id, read_json, safe_name, write_json

EDIT_DIR = UPL / "edit"
_BAD_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def xml_safe(s):
    """Word-ը (XML) չի ընդունում կառավարման նշանները (PDF-ից հանված տեքստում հաճախ կան)."""
    return _BAD_XML.sub("", str(s or ""))
SUPPORTED = (".docx", ".xlsx", ".xlsm", ".csv", ".txt", ".pdf", ".jpg", ".jpeg", ".png", ".webp", ".bmp",
             ".gif", ".tif", ".tiff")
MAX_ROWS, MAX_COLS = 1500, 60


# ================================================================== աշխատանքային պատճեն
def open_copy(src, original_name, owner):
    """Ֆայլը պատճենում է խմբագրման թղթապանակ. -> token."""
    token = new_id("e")
    folder = EDIT_DIR / token
    folder.mkdir(parents=True, exist_ok=True)
    name = safe_name(original_name or Path(src).name, 110, "file")
    dst = folder / name
    shutil.copy2(src, dst)
    write_json(folder / "info.json", dict(owner=owner, name=name, file=name, created=datetime.now().isoformat(timespec="seconds")))
    return token


def session(token, user):
    folder = EDIT_DIR / safe_name(token, 40, "x")
    info = read_json(folder / "info.json", None)
    if not info:
        raise LookupError("Խմբագրման ֆայլը չգտնվեց, բացեք կրկին / Файл для редактирования не найден, откройте заново")
    if info.get("owner") != user.get("login") and user.get("role") != "admin":
        raise PermissionError("Սա ձեր ֆայլը չէ / Это не ваш файл")
    p = folder / info["file"]
    if not p.exists():
        raise LookupError("Ֆայլը ջնջված է / Файл удалён, откройте заново")
    return folder, p, info


def _switch(folder, info, new_path):
    info["file"] = new_path.name
    write_json(folder / "info.json", info)


# ================================================================== DOCX
def _runs(p):
    try:
        out = []
        for item in p.iter_inner_content():
            out.extend(item.runs if hasattr(item, "runs") else [item])
        return out
    except AttributeError:
        return list(p.runs)


def _ptext(p):
    return "".join(r.text for r in _runs(p))


def set_text(p, new):
    """Պարբերության տեքստը փոխում է՝ պահպանելով run-երի ձևաչափը (փոխվում են միայն փոխված նշանները)."""
    new = xml_safe(new)
    runs = _runs(p)
    old = "".join(r.text for r in runs)
    if old == new:
        return False
    if not runs:
        p.add_run(new)
        return True
    bounds, pos = [], 0
    for r in runs:
        bounds.append((pos, pos + len(r.text)))
        pos += len(r.text)

    def run_at(i):
        for k, (a, b) in enumerate(bounds):
            if a <= i < b:
                return k
        return len(runs) - 1

    out = [""] * len(runs)
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, old, new, autojunk=False).get_opcodes():
        if tag == "equal":
            for k, (a, b) in enumerate(bounds):
                lo, hi = max(a, i1), min(b, i2)
                if lo < hi:
                    out[k] += new[j1 + lo - i1:j1 + hi - i1]
        elif j2 > j1:
            k = run_at(i1) if (tag == "replace" and i1 < len(old)) else (run_at(i1 - 1) if i1 > 0 else 0)
            out[k] += new[j1:j2]
    for r, t in zip(runs, out):
        if r.text != t:
            r.text = t
    return True


def _walk(doc):
    """Մարմնի տարրերը հերթականությամբ՝ ('p', Paragraph) կամ ('t', Table)."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    body = doc.element.body
    for el in body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield "p", Paragraph(el, doc)
        elif tag == "tbl":
            yield "t", Table(el, doc)


def _cells(table):
    """Աղյուսակի եզակի վանդակները տողերով (միաձուլված վանդակները՝ մեկ անգամ)."""
    seen, rows = set(), []
    for r in table.rows:
        row = []
        try:
            cells = r.cells
        except Exception:  # noqa — խիստ անկանոն աղյուսակ
            cells = []
        for c in cells:
            if id(c._tc) in seen:
                continue
            seen.add(id(c._tc))
            row.append(c)
        rows.append(row)
    return rows


def _cell_text(c):
    return "\n".join(_ptext(p) for p in c.paragraphs)


def _pstyle(p):
    name = (p.style.name if p.style is not None else "") or ""
    runs = _runs(p)
    bold = bool(runs) and all(r.bold for r in runs if r.text.strip()) and any(r.text.strip() for r in runs)
    if name.lower().startswith(("heading", "title")) or "заголов" in name.lower():
        return "h"
    return "b" if bold else ""


def _align(p):
    try:
        a = p.alignment
        return {1: "center", 2: "right", 3: "justify"}.get(int(a) if a is not None else 0, "")
    except (TypeError, ValueError):
        return ""


def docx_model(path):
    import docx
    d = docx.Document(str(path))
    blocks = []
    for i, (kind, el) in enumerate(_walk(d)):
        if kind == "p":
            blocks.append(dict(t="p", id=i, text=_ptext(el), style=_pstyle(el), align=_align(el)))
        else:
            blocks.append(dict(t="table", id=i, rows=[[_cell_text(c) for c in row] for row in _cells(el)]))
    return dict(type="docx", blocks=blocks, paragraphs=sum(1 for b in blocks if b["t"] == "p"),
                tables=sum(1 for b in blocks if b["t"] == "table"))


def _set_cell(cell, text):
    pars = list(cell.paragraphs)
    lines = str(text).split("\n")
    if "\n".join(_ptext(p) for p in pars) == text:
        return
    for i, line in enumerate(lines):
        if i < len(pars):
            set_text(pars[i], line)
        else:
            el = copy.deepcopy(pars[-1]._p)
            pars[-1]._p.addnext(el)
            from docx.text.paragraph import Paragraph
            np_ = Paragraph(el, pars[-1]._parent)
            for r in _runs(np_)[1:]:
                r.text = ""
            set_text(np_, line)
            pars.append(np_)
    for p in pars[len(lines):]:
        if len(cell.paragraphs) > 1:
            p._p.getparent().remove(p._p)


def docx_save(path, blocks, out_path):
    """blocks՝ [{id, text}] կամ [{id, rows}] — միայն փոխվածները. -> քանի փոփոխություն."""
    import docx
    d = docx.Document(str(path))
    by_id = {int(b["id"]): b for b in blocks or [] if isinstance(b, dict) and str(b.get("id", "")).lstrip("-").isdigit()}
    changed = 0
    for i, (kind, el) in enumerate(_walk(d)):
        b = by_id.get(i)
        if not b:
            continue
        if kind == "p" and "text" in b:
            changed += bool(set_text(el, str(b["text"])))
        elif kind == "t" and isinstance(b.get("rows"), list):
            grid = _cells(el)
            for r, row in enumerate(b["rows"]):
                if r >= len(grid) or not isinstance(row, list):
                    continue
                for c, val in enumerate(row):
                    if c < len(grid[r]) and val is not None and _cell_text(grid[r][c]) != str(val):
                        _set_cell(grid[r][c], str(val))
                        changed += 1
    docmeta.clear_docx(d)   # ձեռքով փոխված՝ ձևի հին տվյալները այլևս ճիշտ չեն
    d.save(str(out_path))
    return changed


def text_to_docx(paragraphs, out_path, title=""):
    import docx
    from docx.shared import Pt
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Sylfaen"
    st.font.size = Pt(11)
    if title:
        d.add_heading(xml_safe(title), level=1)
    for line in paragraphs:
        d.add_paragraph(xml_safe(line))
    d.save(str(out_path))
    return out_path


# ================================================================== XLSX / CSV
def _show(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y %H:%M") if (v.hour or v.minute) else v.strftime("%d.%m.%Y")
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _csv_rows(path):
    import act
    return [[("" if c is None else str(c)) for c in r] for r in act.read_rows(path)]


def sheet_model(path):
    ext = Path(path).suffix.lower()
    if ext == ".csv":
        rows = _csv_rows(path)
        return dict(type="sheet", sheets=[dict(name=Path(path).stem, rows=[r[:MAX_COLS] for r in rows[:MAX_ROWS]],
                                               total_rows=len(rows))])
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, MAX_ROWS),
                              max_col=min(ws.max_column or 1, MAX_COLS), values_only=True):
            rows.append([_show(v) for v in r])
        while rows and not any(rows[-1]):
            rows.pop()
        sheets.append(dict(name=ws.title, rows=rows, total_rows=ws.max_row or 0))
    return dict(type="sheet", sheets=sheets)


def _value(s):
    s = str(s)
    t = s.strip()
    if t.startswith("="):
        return t
    if re.fullmatch(r"-?\d{1,15}", t):
        return int(t)
    if re.fullmatch(r"-?\d+[.,]\d+", t):
        return float(t.replace(",", "."))
    m = re.fullmatch(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", t)
    if m:
        try:
            return datetime(int(m[3]), int(m[2]), int(m[1]))
        except ValueError:
            pass
    return s


def sheet_save(path, changes, out_path):
    ext = Path(path).suffix.lower()
    ch = [c for c in changes or [] if isinstance(c, dict)]
    if ext == ".csv":
        rows = _csv_rows(path)
        for c in ch:
            r, col = int(c.get("r", -1)), int(c.get("c", -1))
            if r < 0 or col < 0:
                continue
            while len(rows) <= r:
                rows.append([])
            while len(rows[r]) <= col:
                rows[r].append("")
            rows[r][col] = str(c.get("v", ""))
        raw = Path(path).read_bytes()[:4000].decode("utf-8", "replace")
        delim = ";" if raw.count(";") > raw.count(",") else ("\t" if raw.count("\t") > raw.count(",") else ",")
        with open(out_path, "w", encoding="utf-8-sig", newline="") as f:   # newline=""՝ առանց կրկնակի \r Windows-ում
            csv.writer(f, delimiter=delim).writerows(rows)
        return len(ch)
    import openpyxl
    wb = openpyxl.load_workbook(str(path), keep_vba=ext == ".xlsm")
    for c in ch:
        s, r, col = int(c.get("s", 0)), int(c.get("r", -1)), int(c.get("c", -1))
        if not (0 <= s < len(wb.worksheets)) or r < 0 or col < 0:
            continue
        cell = wb.worksheets[s].cell(row=r + 1, column=col + 1)
        try:
            cell.value = _value(xml_safe(c.get("v", "")))
        except (AttributeError, ValueError):   # միաձուլված վանդակ (MergedCell) — չի փոխվում
            continue
    wb.save(str(out_path))
    return len(ch)


# ================================================================== բացել
def _pdf_text(path):
    from pypdf import PdfReader
    pars = []
    for page in PdfReader(str(path)).pages[:200]:
        txt = page.extract_text() or ""
        for line in txt.splitlines():
            line = line.strip()
            if line:
                pars.append(line)
    return pars


def model(token, user, convert_scan=None):
    """Ֆայլի տեսքը խմբագրիչի համար. PDF/նկար -> DOCX (մեկ անգամ)."""
    folder, p, info = session(token, user)
    ext = p.suffix.lower()
    meta = docmeta.read(p) if ext in (".docx", ".pdf") else None
    note = ""
    if ext == ".pdf":
        if not meta:
            meta = _plan_from_pdf(p)
        pars = _pdf_text(p)
        if not pars and convert_scan:
            pars = convert_scan(p)
            note = "scan"
        if not pars:
            return dict(type="empty", name=info["name"], meta=meta, token=token,
                        note="no_text")
        out = p.with_suffix(".docx")
        text_to_docx(pars, out)
        _switch(folder, info, out)
        p, ext, note = out, ".docx", note or "pdf"
    elif ext in (".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff"):
        if not convert_scan:
            raise ValueError("Նկարի ճանաչումը միացված չէ / Распознавание фото не настроено (⚙️ Настройки)")
        pars = convert_scan(p)
        if not pars:
            raise ValueError("Նկարում տեքստ չգտնվեց / На фото не найден текст")
        out = p.with_suffix(".docx")
        text_to_docx(pars, out)
        _switch(folder, info, out)
        p, ext, note = out, ".docx", "scan"
    if ext == ".docx":
        res = docx_model(p)
    elif ext in (".xlsx", ".xlsm", ".csv"):
        res = sheet_model(p)
    elif ext == ".txt":
        raw = p.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1251", "replace")
        res = dict(type="text", text=text)
    else:
        raise ValueError(f"«{ext}» ձևաչափը չի խմբագրվում / Формат «{ext}» нельзя редактировать")
    return dict(res, name=info["name"], file=p.name, ext=ext.lstrip("."), meta=meta, token=token, note=note)


def _plan_from_pdf(p):
    """Հին մեդիա պլանի PDF (առանց docmeta)՝ փորձում ենք կարդալ mpimport-ով."""
    try:
        import mpimport
        r = mpimport.parse(p)
    except Exception:  # noqa
        return None
    if r.get("addresses") and r.get("start"):
        return dict(kind="plan", data=dict(client=r["client"], start=r["start"], end=r["end"], slots=r["slots"],
                                           addresses=r["addresses"], clips=[]), source="pdf")
    return None


def save(token, user, payload):
    """-> (out_path, changes)."""
    folder, p, info = session(token, user)
    ext = p.suffix.lower()
    stem = safe_name(Path(info["name"]).stem, 60, "document")
    stem = re.sub(r"(_edited_\d{8}_\d{4})+$", "", stem)
    name = safe_name(payload.get("name") or "", 80, "") or f"{stem}_edited_{datetime.now():%d%m%Y_%H%M}"
    OUT.mkdir(parents=True, exist_ok=True)
    if ext == ".docx":
        out = OUT / f"{Path(name).stem}.docx"
        n = docx_save(p, payload.get("blocks"), out)
    elif ext in (".xlsx", ".xlsm", ".csv"):
        out = OUT / f"{Path(name).stem}{ext}"
        n = sheet_save(p, payload.get("changes"), out)
    elif ext == ".txt":
        out = OUT / f"{Path(name).stem}.txt"
        out.write_text(str(payload.get("text") or ""), encoding="utf-8")
        n = 1
    else:
        raise ValueError("Այս ֆայլը չի պահվում / Этот файл нельзя сохранить")
    # խմբագրումը շարունակվում է արդեն նոր տարբերակից
    shutil.copy2(out, p)
    return out, n


def cleanup(days=3):
    import time
    if not EDIT_DIR.exists():
        return
    cutoff = time.time() - days * 86400
    for d in EDIT_DIR.iterdir():
        try:
            if d.is_dir() and d.stat().st_mtime < cutoff:
                shutil.rmtree(d, ignore_errors=True)
        except OSError:
            continue


def dump(obj):
    return json.dumps(obj, ensure_ascii=False)
