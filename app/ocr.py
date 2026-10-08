# -*- coding: utf-8 -*-
"""🔎 Տեղային OCR / Локальное распознавание текста — без интернета и без ключей.

Նկար / սկան / PDF / Word / Excel -> տեքստ. Նկարները կարդում է Tesseract-ը (անվճար, աշխատում է այս
համակարգչում)՝ հայերեն + ռուսերեն + անգլերեն: Տեղադրում՝ install_ocr.bat (մեկ անգամ):

Лёгкие файлы (Word, Excel, PDF с текстом) читаются напрямую, без OCR."""
import io
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from config import DATA

TESSDATA = DATA / "tessdata"              # install_ocr.bat կլցնի hye/rus/eng լեզուները այստեղ
WANT_LANGS = ("hye", "rus", "eng")
IMAGE_EXT = (".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp", ".tif", ".tiff")
TEXT_EXT = (".pdf", ".docx", ".xlsx", ".xlsm", ".csv", ".txt")
MAX_PDF_PAGES = 15
_TIMEOUT = 120


# ================================================================== Tesseract
def tesseract_cmd():
    cands = [os.getenv("TESSERACT_CMD", ""), shutil.which("tesseract") or "",
             r"C:\Program Files\Tesseract-OCR\tesseract.exe",
             r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
             str(Path(os.getenv("LOCALAPPDATA", "")) / "Programs" / "Tesseract-OCR" / "tesseract.exe")]
    return next((c for c in cands if c and Path(c).is_file()), "")


def _tessdata_dir(cmd):
    """Մեր data/tessdata (եթե լեզուները այնտեղ են), հակառակ դեպքում՝ Tesseract-ի սեփականը."""
    if (TESSDATA / "eng.traineddata").exists():
        return TESSDATA
    own = Path(cmd).parent / "tessdata" if cmd else None
    return own if own and own.is_dir() else None


def languages():
    cmd = tesseract_cmd()
    d = _tessdata_dir(cmd)
    if not cmd or not d:
        return []
    return [l for l in WANT_LANGS if (d / f"{l}.traineddata").exists()]


def status():
    cmd = tesseract_cmd()
    langs = languages()
    return dict(ready=bool(cmd and langs), tesseract=bool(cmd), path=cmd, languages=langs,
                missing=[l for l in WANT_LANGS if l not in langs])


def _prepare(data: bytes):
    """Նկարի պատրաստում OCR-ի համար. ուղղահայաց շրջում (EXIF), մոխրագույն, փոքրերը՝ մեծացնել."""
    from PIL import Image, ImageFilter, ImageOps
    try:
        im = Image.open(io.BytesIO(data))
        im.load()
    except Exception as e:  # noqa
        raise ValueError("Նկարը չբացվեց (HEIC-ը պահեք JPG) / Не удалось открыть фото (HEIC сохраните как JPG)") from e
    im = ImageOps.exif_transpose(im)
    if im.mode in ("RGBA", "LA", "P"):
        bg = Image.new("RGB", im.size, "white")
        im = im.convert("RGBA")
        bg.paste(im, mask=im.split()[-1])
        im = bg
    im = im.convert("L")
    w, h = im.size
    longest = max(w, h)
    if longest < 1800:                     # փոքր նկարներում տառերը մանր են՝ մեծացնում ենք
        k = 1800 / longest
        im = im.resize((int(w * k), int(h * k)), Image.LANCZOS)
    elif longest > 4200:
        k = 4200 / longest
        im = im.resize((int(w * k), int(h * k)), Image.LANCZOS)
    im = ImageOps.autocontrast(im, cutoff=1)
    im = im.filter(ImageFilter.SHARPEN)
    return im


def image_text(data: bytes, psm=3):
    st = status()
    if not st["tesseract"]:
        raise RuntimeError("Tesseract OCR տեղադրված չէ / Не установлен Tesseract OCR — запустите install_ocr.bat "
                           "(один раз, бесплатно, работает без интернета)")
    if not st["languages"]:
        raise RuntimeError("OCR լեզուները չկան / Нет языков OCR (hye, rus, eng) — запустите install_ocr.bat")
    im = _prepare(data)
    cmd = st["path"]
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "in.png"
        im.save(src, "PNG")
        args = [cmd, str(src), "stdout", "-l", "+".join(st["languages"]), "--psm", str(psm), "--oem", "1"]
        td = _tessdata_dir(cmd)
        if td:
            args += ["--tessdata-dir", str(td)]
        flags = 0x08000000 if os.name == "nt" else 0          # CREATE_NO_WINDOW՝ սև պատուհան չբացվի
        try:
            r = subprocess.run(args, capture_output=True, timeout=_TIMEOUT, creationflags=flags)
        except subprocess.TimeoutExpired as e:
            raise RuntimeError("OCR-ը շատ երկար տևեց / Распознавание заняло слишком долго — уменьшите фото") from e
    if r.returncode != 0:
        err = r.stderr.decode("utf-8", "replace").strip().splitlines()
        raise RuntimeError("OCR սխալ / Ошибка OCR: " + (err[-1] if err else f"код {r.returncode}"))
    return r.stdout.decode("utf-8", "replace")


# ================================================================== ֆայլեր
def _decode(raw):
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def docx_text(path):
    import docx
    d = docx.Document(str(path))
    out = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            cells = []
            for c in r.cells:
                tx = c.text.strip()
                if not cells or cells[-1] != tx:
                    cells.append(tx)
            out.append(" | ".join(cells))
    return "\n".join(out)


def sheet_text(path):
    import act
    rows = act.read_rows(path)
    return "\n".join(" | ".join("" if c is None else str(c) for c in r)
                     for r in rows[:4000] if any(c not in (None, "") for c in r))


def pdf_text(path):
    """PDF՝ տեքստային շերտից, իսկ սկանավորված էջերը՝ OCR (էջի նկարներից)."""
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    out = []
    for i, page in enumerate(reader.pages):
        if i >= MAX_PDF_PAGES:
            break
        txt = (page.extract_text() or "").strip()
        if len(txt) >= 40:
            out.append(txt)
            continue
        for img in list(page.images)[:4]:     # սկան՝ էջը մեկ մեծ նկար է
            try:
                if len(img.data) > 20_000:
                    out.append(image_text(img.data))
            except ValueError:
                continue
    return "\n".join(out)


def file_text(path):
    """Ցանկացած աջակցվող ֆայլ -> տեքստ (միայն տեղային միջոցներով)."""
    p = Path(path)
    ext = p.suffix.lower()
    if ext in IMAGE_EXT:
        text = image_text(p.read_bytes())
    elif ext == ".pdf":
        text = pdf_text(p)
    elif ext == ".docx":
        text = docx_text(p)
    elif ext in (".xlsx", ".xlsm", ".csv"):
        text = sheet_text(p)
    elif ext == ".txt":
        text = _decode(p.read_bytes())
    else:
        raise ValueError(f"«{ext}» ձևաչափը չի աջակցվում / Формат «{ext}» не поддерживается")
    return text.replace("\r", "")
