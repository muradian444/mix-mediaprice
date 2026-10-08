# -*- coding: utf-8 -*-
"""📈 Եռամսյակային հաշվետվություն՝ սլայդներ (16:9) -> PDF և ԽՄԲԱԳՐԵԼԻ PPTX.

Օգտատերը ինքն է հավաքում սլայդները (վերնագիր, տեքստ, կետեր, թվեր, գծապատկեր, աղյուսակ,
նկար, ֆայլերի ցանկ) և կարող է ավելացնել իր ֆայլերը (նկարները մտնում են սլայդ, մնացածը՝ ցանկ)."""
import re
from datetime import date, datetime
from pathlib import Path

from reportlab.lib.colors import HexColor, white
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

import mediaplan as mp
from config import ASSETS, DATA, QDIR, UPL
from store import new_id, read_json, safe_name, write_json

FILE = DATA / "quarterly.json"
LOGO = ASSETS / "logo.png"

NAVY, VIOLET, BLUE = HexColor("#1D1B5E"), HexColor("#6A3DE8"), HexColor("#2E8CF0")
INK, MUTED = HexColor("#12152E"), HexColor("#6B7090")
SOFT, TINT, LINE = HexColor("#F5F3FF"), HexColor("#EDE8FE"), HexColor("#DDE0EE")

W, H = 338.667 * mm, 190.5 * mm          # 16:9 (13.333 × 7.5 դյույմ)
M = 20 * mm
CW = W - 2 * M

SLIDE_TYPES = ["cover", "text", "bullets", "metrics", "chart", "table", "image", "photo_text", "gallery", "files",
               "closing"]
TYPE_NAMES = dict(cover="Շապիկ", text="Տեքստ", bullets="Կետեր", metrics="Թվեր (KPI)",
                  chart="Գծապատկեր", table="Աղյուսակ", image="Նկար", photo_text="Տեքստ + նկար",
                  gallery="Ֆոտոշարք", files="Ֆայլեր", closing="Եզրափակիչ")
QUARTERS = {1: "I եռամսյակ", 2: "II եռամսյակ", 3: "III եռամսյակ", 4: "IV եռամսյակ"}


# ================================================================== մոդել
def _all():
    d = read_json(FILE, [])
    return d if isinstance(d, list) else []


def _num(v, default=0.0):
    try:
        return float(str(v).replace(" ", "").replace(",", "."))
    except (TypeError, ValueError):
        return default


def _norm_slide(s):
    s = s if isinstance(s, dict) else {}
    t = s.get("type") if s.get("type") in SLIDE_TYPES else "text"
    out = dict(id=str(s.get("id") or new_id("s")), type=t,
               title=str(s.get("title") or "").strip(), subtitle=str(s.get("subtitle") or "").strip(),
               body=str(s.get("body") or "").strip(), caption=str(s.get("caption") or "").strip(),
               items=[], image=str(s.get("image") or ""), unit=str(s.get("unit") or "").strip(),
               head=[str(h) for h in (s.get("head") or [])], rows=[], images=[])
    if t in ("bullets", "photo_text"):
        out["items"] = [str(x).strip() for x in (s.get("items") or []) if str(x).strip()]
    elif t in ("metrics", "chart"):
        for it in (s.get("items") or []):
            it = it if isinstance(it, dict) else {}
            label = str(it.get("label") or "").strip()
            if not label and not str(it.get("value") or "").strip():
                continue
            out["items"].append(dict(label=label, value=str(it.get("value") or "").strip(),
                                     note=str(it.get("note") or "").strip()))
    elif t == "table":
        for r in (s.get("rows") or []):
            cells = [str(c) for c in (r if isinstance(r, list) else [r])]
            if any(c.strip() for c in cells):
                out["rows"].append(cells)
    elif t == "files":
        out["items"] = [str(x).strip() for x in (s.get("items") or []) if str(x).strip()]
    elif t == "gallery":
        for im in (s.get("images") or [])[:6]:
            im = im if isinstance(im, dict) else dict(path=str(im))
            if im.get("path"):
                out["images"].append(dict(path=str(im["path"]), caption=str(im.get("caption") or "").strip()[:120]))
    return out


def _norm(d):
    d = d if isinstance(d, dict) else {}
    now = datetime.now().isoformat(timespec="seconds")
    q = d.get("quarter")
    try:
        q = int(q)
        q = q if q in QUARTERS else (date.today().month - 1) // 3 + 1
    except (TypeError, ValueError):
        q = (date.today().month - 1) // 3 + 1
    return dict(id=str(d.get("id") or new_id("q")),
                title=str(d.get("title") or "Եռամսյակային հաշվետվություն").strip(),
                client=str(d.get("client") or "").strip(),
                quarter=q, year=int(_num(d.get("year"), date.today().year)) or date.today().year,
                author=str(d.get("author") or "").strip(),
                summary=str(d.get("summary") or "").strip(),
                prompt=str(d.get("prompt") or "")[:20000],
                slides=[_norm_slide(s) for s in (d.get("slides") or [])],
                attachments=[a for a in (d.get("attachments") or []) if isinstance(a, dict)],
                owner=str(d.get("owner") or ""),
                created=d.get("created") or now, updated=d.get("updated") or now)


def visible(deck, user):
    """Հաշվետվությունը տեսնում է իր հեղինակը և ադմինը (հին՝ առանց հեղինակի՝ միայն ադմինը)."""
    if not deck:
        return False
    if not user or user.get("role") == "admin":
        return True
    return deck.get("owner") == user.get("login")


def decks(user=None):
    out = [_norm(d) for d in _all()]
    out = [d for d in out if visible(d, user)]
    out.sort(key=lambda d: d["updated"], reverse=True)
    return [dict(d, slides=len(d["slides"])) for d in out]


def get(deck_id):
    for d in _all():
        if str(d.get("id")) == str(deck_id):
            return _norm(d)
    return None


def create(data, owner=""):
    d = _norm(dict(data or {}, id=new_id("q"), owner=owner))
    if not d["slides"]:
        d["slides"] = [_norm_slide(dict(type="cover", title=d["title"],
                                        subtitle=f"{QUARTERS[d['quarter']]} {d['year']}")),
                       _norm_slide(dict(type="metrics", title="Հիմնական թվերը", items=[
                           dict(label="Հեռարձակումներ", value="0", note="սփոթ"),
                           dict(label="Հասցեներ", value="0", note="օբյեկտ"),
                           dict(label="Կատարում", value="0%", note="պլանից")])),
                       _norm_slide(dict(type="bullets", title="Արդյունքներ", items=["", "", ""]))]
    rows = _all()
    rows.append(d)
    write_json(FILE, rows)
    return d


def update(deck_id, data):
    rows = _all()
    for i, raw in enumerate(rows):
        if str(raw.get("id")) != str(deck_id):
            continue
        old = _norm(raw)
        merged = _norm({**old, **{k: v for k, v in (data or {}).items() if k not in ("id", "created", "owner")}})
        merged["id"], merged["created"], merged["owner"] = old["id"], old["created"], old["owner"]
        merged["updated"] = datetime.now().isoformat(timespec="seconds")
        rows[i] = merged
        write_json(FILE, rows)
        return merged
    raise ValueError("Հաշվետվությունը չգտնվեց")


def delete(deck_id):
    rows = _all()
    keep = [r for r in rows if str(r.get("id")) != str(deck_id)]
    if len(keep) == len(rows):
        raise ValueError("Հաշվետվությունը չգտնվեց")
    write_json(FILE, keep)
    folder = QDIR / str(deck_id)
    if folder.exists():
        import shutil
        shutil.rmtree(folder, ignore_errors=True)
    return True


def media_dir(deck_id):
    p = QDIR / safe_name(deck_id, 40, "deck")
    p.mkdir(parents=True, exist_ok=True)
    return p


def attach(deck_id, filename, data: bytes):
    """Օգտատիրոջ ֆայլը՝ նկարները դառնում են սլայդ, մնացածը՝ «Ֆայլեր» ցանկ."""
    d = get(deck_id)
    if not d:
        raise ValueError("Հաշվետվությունը չգտնվեց")
    folder = media_dir(deck_id)
    name = safe_name(filename, 80, "file")
    p = folder / name
    for i in range(2, 99):
        if not p.exists():
            break
        p = folder / f"{Path(name).stem} ({i}){Path(name).suffix}"
    p.write_bytes(data)
    rel = f"{p.parent.name}/{p.name}"
    is_img = p.suffix.lower() in (".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp")
    att = dict(name=p.name, path=rel, size=len(data), kind="image" if is_img else "file")
    update(deck_id, dict(attachments=d["attachments"] + [att]))
    return att


def attachment_path(rel):
    """data/quarterly/<deck>/<file> -> բացարձակ ուղի (ստուգված)."""
    rel = str(rel or "").replace("\\", "/")
    if ".." in rel.split("/"):
        raise ValueError("Ուղին սխալ է")
    p = (QDIR / rel).resolve()
    if QDIR.resolve() not in p.parents:
        raise ValueError("Ուղին սխալ է")
    if not p.exists():
        raise ValueError("Ֆայլը չգտնվեց")
    return p


def import_numbers(path):
    """xlsx/csv -> պատրաստի «Թվեր» սլայդ (օգտագործում է ԱԿՏ-ի նույն կարդացողը)."""
    import act
    rows = act.read_rows(path)
    dates = []
    for r in rows:
        for c in r:
            d = act._to_date(c)
            if d:
                dates.append(d)
                break
    if not dates:
        raise ValueError("Ֆայլում ամսաթվեր չգտնվեցին")
    start, end = min(dates), max(dates)
    days, clips = act.aggregate(rows, start, end)
    planned, played, missed = act.totals(days)
    items = [dict(label="Նախատեսված", value=f"{planned:,}".replace(",", " "), note="սփոթ"),
             dict(label="Հեռարձակված", value=f"{played:,}".replace(",", " "), note="սփոթ"),
             dict(label="Չհեռարձակված", value=f"{missed:,}".replace(",", " "), note="սփոթ"),
             dict(label="Կատարում", value=f"{(100 * played / planned):.1f}%" if planned else "—", note="պլանից")]
    chart = [dict(label=f"{d:%d.%m}", value=str(v["played"])) for d, v in list(days.items())[:31]]
    return dict(period=f"{start:%d.%m.%Y} — {end:%d.%m.%Y}", metrics=items, chart=chart,
                clips=[dict(label=k, value=str(v)) for k, v in list(clips.items())[:12]])


# ================================================================== PDF
class Deck:
    def __init__(self, path, title):
        mp._fonts()
        self.c = canvas.Canvas(str(path), pagesize=(W, H))
        self.c.setTitle(title)
        self.c.setAuthor("Mix Media production")
        self.n = 0

    # --- գծագրման գործիքներ
    def text(self, x, y, s, size=12, bold=False, color=INK, align="l"):
        c = self.c
        c.setFont("MM-B" if bold else "MM", size)
        c.setFillColor(color)
        {"l": c.drawString, "r": c.drawRightString, "c": c.drawCentredString}[align](x, y, str(s))

    def wrap(self, s, size, width, bold=False):
        font = "MM-B" if bold else "MM"
        out, cur = [], ""
        for word in str(s).split():
            t = (cur + " " + word).strip()
            if stringWidth(t, font, size) <= width or not cur:
                cur = t
            else:
                out.append(cur)
                cur = word
        if cur:
            out.append(cur)
        return out or [""]

    def para(self, x, y, s, size, width, bold=False, color=INK, lead=1.45, max_lines=None):
        lines = self.wrap(s, size, width, bold)
        if max_lines and len(lines) > max_lines:
            lines = lines[:max_lines]
            lines[-1] = lines[-1][:max(0, len(lines[-1]) - 1)] + "…"
        for ln in lines:
            self.text(x, y, ln, size, bold, color)
            y -= size * lead
        return y

    def rrect(self, x, y, w, h, r=3 * mm, fill=white, line=None, lw=0.7):
        c = self.c
        if fill is not None:
            c.setFillColor(fill)
        if line is not None:
            c.setStrokeColor(line)
            c.setLineWidth(lw)
        c.roundRect(x, y, w, h, r, fill=1 if fill is not None else 0, stroke=1 if line is not None else 0)

    def grad(self, x, y, w, h, radius=0, c0=BLUE, c1=VIOLET):
        c = self.c
        c.saveState()
        p = c.beginPath()
        p.roundRect(x, y, w, h, radius) if radius else p.rect(x, y, w, h)
        c.clipPath(p, stroke=0, fill=0)
        c.linearGradient(x, y, x + w, y + h, (c0, c1))
        c.restoreState()

    def logo(self, x, y_top, w, dark=False):
        if not LOGO.exists():
            self.text(x, y_top - 6 * mm, "MIXMEDIA", 13, True, white if dark else NAVY)
            return 8 * mm
        from PIL import Image
        iw, ih = Image.open(LOGO).size
        h = w * ih / iw
        self.c.drawImage(str(LOGO), x, y_top - h, w, h, mask="auto")
        return h

    # --- սլայդի կարկաս
    def start(self, title="", kicker="", dark=False):
        self.n += 1
        if dark:
            self.c.setFillColor(NAVY)
            self.c.rect(0, 0, W, H, fill=1, stroke=0)
            self.grad(0, H - 2.2 * mm, W, 2.2 * mm)
        else:
            self.c.setFillColor(white)
            self.c.rect(0, 0, W, H, fill=1, stroke=0)
            self.grad(0, H - 2.2 * mm, W, 2.2 * mm)
        y = H - 18 * mm
        if kicker:
            self.text(M, y, kicker.upper(), 8, True, BLUE if not dark else TINT)
            y -= 7 * mm
        if title:
            for ln in self.wrap(title, 22, CW - 40 * mm, True)[:2]:
                self.text(M, y - 7 * mm, ln, 22, True, white if dark else NAVY)
                y -= 11 * mm
            self.grad(M, y - 2 * mm, 26 * mm, 1.3 * mm)
            y -= 10 * mm
        self.y = y

    def footer(self, note=""):
        self.c.setStrokeColor(LINE)
        self.c.setLineWidth(0.6)
        self.c.line(M, 12 * mm, W - M, 12 * mm)
        self.text(M, 7.5 * mm, note or f"{mp.COMPANY_LINE}", 7.2, False, MUTED)
        self.text(W - M, 7.5 * mm, str(self.n), 8, True, VIOLET, "r")

    def end(self, dark=False, note=""):
        if not dark:
            self.footer(note)
        else:
            self.text(W - M, 7.5 * mm, str(self.n), 8, True, TINT, "r")
        self.c.showPage()

    def save(self):
        self.c.save()


def _cover(pg, d, s):
    pg.start(dark=True)
    pg.logo(M, H - 16 * mm, 54 * mm, dark=True)
    y = H * 0.55
    pg.text(M, y + 16 * mm, f"{QUARTERS[d['quarter']]} · {d['year']}", 10, True, HexColor("#9F8CF5"))
    for ln in pg.wrap(s["title"] or d["title"], 34, CW - 60 * mm, True)[:2]:
        pg.text(M, y, ln, 34, True, white)
        y -= 15 * mm
    if s["subtitle"] or d["client"]:
        pg.text(M, y - 1 * mm, s["subtitle"] or d["client"], 13, False, HexColor("#C9C2F2"))
    pg.grad(M, 34 * mm, 40 * mm, 1.6 * mm)
    info = [("ՊԱՏՎԻՐԱՏՈՒ", d["client"] or "—"), ("ԺԱՄԱՆԱԿԱՀԱՏՎԱԾ", f"{QUARTERS[d['quarter']]} {d['year']}"),
            ("ԿԱԶՄԵԼ Է", d["author"] or mp.COMPANY_LINE.split(",")[0])]
    x = M
    for lab, val in info:
        pg.text(x, 25 * mm, lab, 7, True, HexColor("#8E85C9"))
        pg.text(x, 19 * mm, pg.wrap(val, 11, CW / 3 - 10 * mm, True)[0], 11, True, white)
        x += CW / 3
    pg.end(dark=True)


def _text_slide(pg, d, s):
    pg.start(s["title"], "հաշվետվություն")
    if s["subtitle"]:
        pg.y = pg.para(M, pg.y, s["subtitle"], 12, CW, True, VIOLET) - 2 * mm
    body = s["body"] or ""
    if len(body) <= 700:
        pg.para(M, pg.y, body, 12, CW - 30 * mm, color=INK, max_lines=14)
    else:  # երկար տեքստը՝ երկու սյունակով
        half = (CW - 12 * mm) / 2
        lines = pg.wrap(body, 11.5, half)
        cut = (len(lines) + 1) // 2
        for ci, chunk in enumerate((lines[:cut], lines[cut:cut * 2])):
            y = pg.y
            for ln in chunk:
                pg.text(M + ci * (half + 12 * mm), y, ln, 11.5, False, INK)
                y -= 11.5 * 1.45
    pg.end()


def _bullets(pg, d, s):
    pg.start(s["title"], "հաշվետվություն")
    items = [i for i in s["items"] if i.strip()] or ["—"]
    cols = 2 if len(items) > 6 else 1
    per = (len(items) + cols - 1) // cols
    colw = (CW - 14 * mm) / cols
    for ci in range(cols):
        y = pg.y
        for k, it in enumerate(items[ci * per:(ci + 1) * per], 1):
            x = M + ci * (colw + 14 * mm)
            n = ci * per + k
            pg.grad(x, y - 7.5 * mm, 7.5 * mm, 7.5 * mm, radius=2 * mm)
            pg.text(x + 3.75 * mm, y - 5.2 * mm, f"{n:02d}", 7.5, True, white, "c")
            y = pg.para(x + 12 * mm, y - 4 * mm, it, 12, colw - 14 * mm, max_lines=3) - 4 * mm
    pg.end()


def _metrics(pg, d, s):
    pg.start(s["title"], "հիմնական ցուցանիշներ")
    items = s["items"][:8] or [dict(label="—", value="0", note="")]
    cols = min(len(items), 4)
    rows = (len(items) + cols - 1) // cols
    gap = 7 * mm
    cw = (CW - gap * (cols - 1)) / cols
    ch = min(42 * mm, (pg.y - 22 * mm) / rows - gap + gap / rows)
    for i, it in enumerate(items):
        r, c = divmod(i, cols)
        x = M + c * (cw + gap)
        y = pg.y - r * (ch + gap)
        pg.rrect(x, y - ch, cw, ch, r=3.5 * mm, fill=SOFT, line=None)
        pg.grad(x, y - ch, 1.8 * mm, ch)
        pg.text(x + 8 * mm, y - 9 * mm, pg.wrap(it["label"], 9, cw - 14 * mm, True)[0], 9, True, MUTED)
        val = pg.wrap(it["value"] or "—", 26, cw - 14 * mm, True)[0]
        size = 26
        while stringWidth(val, "MM-B", size) > cw - 14 * mm and size > 12:
            size -= 1
        pg.text(x + 8 * mm, y - 24 * mm, val, size, True, NAVY)
        if it["note"]:
            pg.text(x + 8 * mm, y - ch + 6 * mm, pg.wrap(it["note"], 8.5, cw - 14 * mm)[0], 8.5, False, MUTED)
    pg.end()


def _chart(pg, d, s):
    pg.start(s["title"], "դինամիկա")
    items = [i for i in s["items"] if i["label"]][:24] or [dict(label="—", value="0", note="")]
    vals = [max(0.0, _num(i["value"])) for i in items]
    top = max(vals) or 1.0
    base_y = 28 * mm
    area_h = pg.y - base_y - 8 * mm
    gap = 3.5 * mm if len(items) <= 14 else 2 * mm
    bw = (CW - gap * (len(items) - 1)) / len(items)
    pg.c.setStrokeColor(LINE)
    pg.c.setLineWidth(0.6)
    for k in range(5):
        gy = base_y + area_h * k / 4
        pg.c.line(M, gy, W - M, gy)
        pg.text(M - 3 * mm, gy - 1 * mm, f"{top * k / 4:,.0f}".replace(",", " "), 7, False, MUTED, "r")
    for i, (it, v) in enumerate(zip(items, vals)):
        x = M + i * (bw + gap)
        h = max(0.6 * mm, area_h * v / top)
        pg.grad(x, base_y, bw, h, radius=min(1.6 * mm, bw / 2), c0=BLUE, c1=VIOLET)
        lab = pg.wrap(it["label"], 7.2, bw + gap, True)[0]
        pg.text(x + bw / 2, base_y - 5 * mm, lab, 7.2, False, MUTED, "c")
        if len(items) <= 16:
            pg.text(x + bw / 2, base_y + h + 2 * mm, f"{v:,.0f}".replace(",", " "), 7.5, True, NAVY, "c")
    if s["unit"]:
        pg.text(W - M, base_y - 11 * mm, s["unit"], 8, False, MUTED, "r")
    pg.end()


def _table(pg, d, s):
    pg.start(s["title"], "աղյուսակ")
    head = s["head"] or (s["rows"][0] if s["rows"] else ["—"])
    rows = s["rows"] if s["head"] else s["rows"][1:]
    ncol = max(len(head), max((len(r) for r in rows), default=1))
    head = (head + [""] * ncol)[:ncol]
    cw = CW / ncol
    y = pg.y
    rh = 9 * mm
    pg.rrect(M, y - rh, CW, rh, r=2 * mm, fill=NAVY, line=None)
    for i, h in enumerate(head):
        pg.text(M + i * cw + 4 * mm, y - 6 * mm, pg.wrap(h, 9, cw - 8 * mm, True)[0], 9, True, white)
    y -= rh
    limit = int((y - 20 * mm) / (7.5 * mm))
    for ri, r in enumerate(rows[:limit]):
        cells = (list(r) + [""] * ncol)[:ncol]
        if ri % 2 == 0:
            pg.rrect(M, y - 7.5 * mm, CW, 7.5 * mm, r=0, fill=SOFT, line=None)
        for i, cell in enumerate(cells):
            pg.text(M + i * cw + 4 * mm, y - 5.2 * mm, pg.wrap(cell, 9, cw - 8 * mm)[0], 9, False, INK)
        y -= 7.5 * mm
        pg.c.setStrokeColor(LINE)
        pg.c.setLineWidth(0.4)
        pg.c.line(M, y, W - M, y)
    if len(rows) > limit:
        pg.text(M, y - 6 * mm, f"… ևս {len(rows) - limit} տող", 8, False, MUTED)
    pg.end()


def _image(pg, d, s):
    pg.start(s["title"], "նկար")
    path = None
    try:
        path = attachment_path(s["image"]) if s["image"] else None
    except ValueError:
        path = None
    box_h = pg.y - 24 * mm
    if path:
        from PIL import Image
        iw, ih = Image.open(path).size
        scale = min(CW / iw, box_h / ih)
        w, h = iw * scale, ih * scale
        pg.c.drawImage(str(path), M + (CW - w) / 2, pg.y - h, w, h, mask="auto")
        if s["caption"]:
            pg.text(W / 2, pg.y - h - 7 * mm, s["caption"], 9, False, MUTED, "c")
    else:
        pg.rrect(M, pg.y - box_h, CW, box_h, r=4 * mm, fill=SOFT, line=LINE)
        pg.text(W / 2, pg.y - box_h / 2, "Նկարը ընտրված չէ (ավելացրեք ֆայլ)", 12, False, MUTED, "c")
    pg.end()


def _img_path(rel):
    try:
        return attachment_path(rel) if rel else None
    except ValueError:
        return None


def _draw_photo(pg, path, x, y, w, h, cover=True, radius=3 * mm):
    """Նկարը տուփի մեջ. cover=True՝ լցնում է ամբողջը (կտրում է եզրերը), կլորացված անկյուններով."""
    from PIL import Image, ImageOps
    im = ImageOps.exif_transpose(Image.open(path))
    iw, ih = im.size
    k = max(w / iw, h / ih) if cover else min(w / iw, h / ih)
    dw, dh = iw * k, ih * k
    c = pg.c
    c.saveState()
    p = c.beginPath()
    p.roundRect(x, y, w, h, radius)
    c.clipPath(p, stroke=0, fill=0)
    from reportlab.lib.utils import ImageReader
    if im.mode not in ("RGB", "L"):
        im = im.convert("RGB")
    im.thumbnail((1600, 1600))
    c.drawImage(ImageReader(im), x + (w - dw) / 2, y + (h - dh) / 2, dw, dh)
    c.restoreState()


def _photo_text(pg, d, s):
    """Տեքստ (կամ կետեր) ձախում + նկար աջում՝ ավտո-պրեզենտացիայի հիմնական սլայդը."""
    pg.start(s["title"], "հաշվետվություն")
    path = _img_path(s["image"])
    tw = CW * 0.5 if path else CW
    box_h = pg.y - 20 * mm
    if path:
        px, pw = M + tw + 10 * mm, CW - tw - 10 * mm
        _draw_photo(pg, path, px, 20 * mm, pw, box_h)
        if s["caption"]:
            pg.rrect(px, 20 * mm, pw, 9 * mm, r=0, fill=HexColor("#1D1B5E"), line=None)
            pg.text(px + 4 * mm, 23.2 * mm, pg.wrap(s["caption"], 8.5, pw - 8 * mm)[0], 8.5, False, white)
    y = pg.y
    items = [i for i in s["items"] if i.strip()]
    if items:
        for k, it in enumerate(items[:7], 1):
            pg.grad(M, y - 6.5 * mm, 6.5 * mm, 6.5 * mm, radius=1.8 * mm)
            pg.text(M + 3.25 * mm, y - 4.5 * mm, f"{k:02d}", 7, True, white, "c")
            y = pg.para(M + 10 * mm, y - 3.6 * mm, it, 11.5, tw - 12 * mm, max_lines=3) - 3.5 * mm
            if y < 24 * mm:
                break
    else:
        pg.para(M, y, s["body"] or "", 12, tw - 4 * mm, color=INK, max_lines=int(box_h / (12 * 1.45)))
    pg.end()


def _gallery(pg, d, s):
    pg.start(s["title"] or "Ֆոտոշարք", "ֆոտոշարք")
    ims = [(im, _img_path(im["path"])) for im in s["images"]]
    ims = [(im, p) for im, p in ims if p][:6]
    if not ims:
        pg.text(M, pg.y - 8 * mm, "Նկարներ չկան", 12, False, MUTED)
        pg.end()
        return
    n = len(ims)
    cols = 1 if n == 1 else 2 if n in (2, 4) else 3
    rows = (n + cols - 1) // cols
    gap = 6 * mm
    top, bottom = pg.y, 18 * mm
    cw = (CW - gap * (cols - 1)) / cols
    ch = (top - bottom - gap * (rows - 1)) / rows
    for i, (im, p) in enumerate(ims):
        r, c = divmod(i, cols)
        x = M + c * (cw + gap)
        y = top - (r + 1) * ch - r * gap
        _draw_photo(pg, p, x, y, cw, ch)
        if im.get("caption"):
            pg.rrect(x, y, cw, 8 * mm, r=0, fill=HexColor("#1D1B5E"), line=None)
            pg.text(x + 3 * mm, y + 2.8 * mm, pg.wrap(im["caption"], 8, cw - 6 * mm)[0], 8, False, white)
    pg.end()


def _files(pg, d, s):
    pg.start(s["title"] or "Կցված ֆայլեր", "հավելված")
    names = s["items"] or [a["name"] for a in d["attachments"]]
    if not names:
        pg.text(M, pg.y - 6 * mm, "Ֆայլեր չկան", 12, False, MUTED)
    y = pg.y
    for i, name in enumerate(names[:16], 1):
        pg.rrect(M, y - 10 * mm, CW, 9 * mm, r=2.5 * mm, fill=SOFT if i % 2 else white, line=LINE)
        pg.text(M + 5 * mm, y - 6.5 * mm, f"{i:02d}", 8.5, True, VIOLET)
        pg.text(M + 16 * mm, y - 6.5 * mm, pg.wrap(name, 10, CW - 30 * mm, True)[0], 10, True, INK)
        y -= 11.5 * mm
    if len(names) > 16:
        pg.text(M, y - 4 * mm, f"… ևս {len(names) - 16} ֆայլ", 8.5, False, MUTED)
    pg.end()


def _closing(pg, d, s):
    pg.start(dark=True)
    pg.logo(M, H - 16 * mm, 50 * mm, dark=True)
    pg.text(M, H * 0.5, s["title"] or "Շնորհակալություն", 32, True, white)
    if s["body"]:
        pg.para(M, H * 0.5 - 14 * mm, s["body"], 12, CW - 80 * mm, color=HexColor("#C9C2F2"), max_lines=4)
    y = 34 * mm
    pg.grad(M, y + 8 * mm, 40 * mm, 1.6 * mm)
    contacts = (("ՀԵՌԱԽՈՍ", mp.COMPANY_LINE.split("Հեռ.՝")[-1].strip()),
                ("ԿԱՅՔ", "mix-media.am"), ("ՀԱՍՑԵ", "ՀՀ, ք. Երևան, Լևոնյան 48"))
    for i, (lab, val) in enumerate(contacts):
        x = M + i * CW / 3
        pg.text(x, y, lab, 7, True, HexColor("#8E85C9"))
        pg.text(x, y - 6 * mm, pg.wrap(val, 11, CW / 3 - 8 * mm, True)[0], 11, True, white)
    pg.end(dark=True)


RENDERERS = dict(cover=_cover, text=_text_slide, bullets=_bullets, metrics=_metrics,
                 chart=_chart, table=_table, image=_image, photo_text=_photo_text, gallery=_gallery,
                 files=_files, closing=_closing)


def render_pdf(deck, out_path):
    d = _norm(deck)
    pg = Deck(out_path, f"{d['title']} — {QUARTERS[d['quarter']]} {d['year']}")
    slides = d["slides"] or [_norm_slide(dict(type="cover", title=d["title"]))]
    for s in slides:
        try:
            RENDERERS.get(s["type"], _text_slide)(pg, d, s)
        except Exception:  # մեկ սլայդի սխալը չի կոտրում ամբողջ ֆայլը
            pg.start(s.get("title") or "Սլայդ", "սխալ")
            pg.text(M, pg.y - 8 * mm, "Այս սլայդը նկարել չհաջողվեց՝ ստուգեք տվյալները", 12, False, MUTED)
            pg.end()
    pg.save()
    return Path(out_path)


# ================================================================== PPTX (խմբագրելի)
def _crop_tmp(path, ratio):
    """Նկարը կտրում ենք տուփի համամասնությամբ (w/h) -> BytesIO (PowerPoint-ի համար)."""
    import io
    from PIL import Image, ImageOps
    im = ImageOps.exif_transpose(Image.open(path)).convert("RGB")
    iw, ih = im.size
    if iw / ih > ratio:
        nw = int(ih * ratio)
        im = im.crop(((iw - nw) // 2, 0, (iw - nw) // 2 + nw, ih))
    else:
        nh = int(iw / ratio)
        im = im.crop((0, (ih - nh) // 2, iw, (ih - nh) // 2 + nh))
    im.thumbnail((1800, 1800))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=86)
    buf.seek(0)
    return buf


def render_pptx(deck, out_path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    d = _norm(deck)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    navy, violet, blue = RGBColor(0x1D, 0x1B, 0x5E), RGBColor(0x6A, 0x3D, 0xE8), RGBColor(0x2E, 0x8C, 0xF0)
    ink, muted, softc = RGBColor(0x12, 0x15, 0x2E), RGBColor(0x6B, 0x70, 0x90), RGBColor(0xF5, 0xF3, 0xFF)
    SW, SH = prs.slide_width, prs.slide_height

    def box(sl, x, y, w, h, text, size=14, bold=False, color=ink, align=PP_ALIGN.LEFT, wrap=True):
        tb = sl.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = wrap
        lines = str(text).split("\n") or [""]
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = ln
            r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
            r.font.name = "Sylfaen"
        return tb

    def rect(sl, x, y, w, h, color, line=None):
        from pptx.enum.shapes import MSO_SHAPE
        sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        sh.fill.solid()
        sh.fill.fore_color.rgb = color
        sh.adjustments[0] = 0.06
        if line:
            sh.line.color.rgb = line
            sh.line.width = Pt(0.75)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        if sh.has_text_frame:
            sh.text_frame.text = ""
        return sh

    def bg(sl, color):
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = color

    def head(sl, title, kicker=""):
        rect(sl, 0, 0, SW, Inches(0.07), violet)
        if kicker:
            box(sl, Inches(0.8), Inches(0.45), Inches(8), Inches(0.3), kicker.upper(), 10, True, blue)
        box(sl, Inches(0.8), Inches(0.75), SW - Inches(1.6), Inches(0.9), title, 28, True, navy)
        rect(sl, Inches(0.8), Inches(1.65), Inches(1.0), Inches(0.05), blue)

    def foot(sl, n):
        box(sl, Inches(0.8), SH - Inches(0.55), Inches(9), Inches(0.3), mp.COMPANY_LINE, 8, False, muted)
        box(sl, SW - Inches(1.3), SH - Inches(0.55), Inches(0.5), Inches(0.3), str(n), 9, True, violet, PP_ALIGN.RIGHT)

    for n, s in enumerate(d["slides"] or [_norm_slide(dict(type="cover", title=d["title"]))], 1):
        sl = prs.slides.add_slide(blank)
        t = s["type"]
        if t in ("cover", "closing"):
            bg(sl, navy)
            rect(sl, 0, 0, SW, Inches(0.07), violet)
            if LOGO.exists():
                sl.shapes.add_picture(str(LOGO), Inches(0.8), Inches(0.6), height=Inches(0.75))
            box(sl, Inches(0.8), Inches(2.3), SW - Inches(2), Inches(0.4),
                f"{QUARTERS[d['quarter']]} · {d['year']}", 12, True, RGBColor(0x9F, 0x8C, 0xF5))
            box(sl, Inches(0.8), Inches(2.8), SW - Inches(2), Inches(1.6),
                s["title"] or (d["title"] if t == "cover" else "Շնորհակալություն"), 40, True, RGBColor(255, 255, 255))
            box(sl, Inches(0.8), Inches(4.5), SW - Inches(2), Inches(0.8),
                s["subtitle"] or s["body"] or d["client"], 14, False, RGBColor(0xC9, 0xC2, 0xF2))
            x = Inches(0.8)
            for lab, val in (("ՊԱՏՎԻՐԱՏՈՒ", d["client"] or "—"),
                             ("ԺԱՄԱՆԱԿԱՀԱՏՎԱԾ", f"{QUARTERS[d['quarter']]} {d['year']}"),
                             ("ԿԱՊ", mp.COMPANY_LINE.split("Հեռ.՝")[-1].strip())):
                box(sl, x, SH - Inches(1.5), Inches(3.6), Inches(0.3), lab, 9, True, RGBColor(0x8E, 0x85, 0xC9))
                box(sl, x, SH - Inches(1.2), Inches(3.6), Inches(0.4), val, 13, True, RGBColor(255, 255, 255))
                x += Inches(3.9)
            continue
        head(sl, s["title"] or TYPE_NAMES.get(t, ""), "հաշվետվություն")
        top = Inches(2.1)
        if t == "text":
            box(sl, Inches(0.8), top, SW - Inches(1.6), SH - top - Inches(1), s["body"], 14)
        elif t == "bullets":
            y = top
            for i, it in enumerate([x for x in s["items"] if x.strip()] or ["—"], 1):
                rect(sl, Inches(0.8), y, Inches(0.33), Inches(0.33), violet)
                box(sl, Inches(0.82), y + Inches(0.03), Inches(0.3), Inches(0.3), f"{i:02d}", 9, True,
                    RGBColor(255, 255, 255), PP_ALIGN.CENTER)
                box(sl, Inches(1.3), y - Inches(0.03), SW - Inches(2.3), Inches(0.5), it, 14)
                y += Inches(0.55)
        elif t == "metrics":
            items = s["items"][:8] or [dict(label="—", value="0", note="")]
            cols = min(len(items), 4)
            cw = (SW - Inches(1.6) - Inches(0.25) * (cols - 1)) / cols
            for i, it in enumerate(items):
                r, c = divmod(i, cols)
                x = Inches(0.8) + c * (cw + Inches(0.25))
                y = top + r * Inches(1.9)
                rect(sl, x, y, cw, Inches(1.6), softc)
                rect(sl, x, y, Inches(0.07), Inches(1.6), violet)
                box(sl, x + Inches(0.3), y + Inches(0.15), cw - Inches(0.5), Inches(0.3), it["label"], 10, True, muted)
                box(sl, x + Inches(0.3), y + Inches(0.5), cw - Inches(0.5), Inches(0.6), it["value"] or "—", 26, True, navy)
                box(sl, x + Inches(0.3), y + Inches(1.15), cw - Inches(0.5), Inches(0.3), it["note"], 9, False, muted)
        elif t == "chart":
            items = [i for i in s["items"] if i["label"]][:20] or [dict(label="—", value="0", note="")]
            vals = [max(0.0, _num(i["value"])) for i in items]
            top_v = max(vals) or 1.0
            area_h, base = Inches(3.6), SH - Inches(1.4)
            bw = (SW - Inches(1.6)) / len(items) * 0.72
            step = (SW - Inches(1.6)) / len(items)
            for i, (it, v) in enumerate(zip(items, vals)):
                h = Emu(int(area_h * (v / top_v))) if v else Inches(0.03)
                x = Inches(0.8) + i * step
                rect(sl, x, base - h, Emu(int(bw)), h, blue if i % 2 else violet)
                box(sl, x - Inches(0.1), base + Inches(0.05), Emu(int(bw)) + Inches(0.2), Inches(0.3),
                    it["label"], 8, False, muted, PP_ALIGN.CENTER)
                box(sl, x - Inches(0.1), base - h - Inches(0.3), Emu(int(bw)) + Inches(0.2), Inches(0.3),
                    f"{v:,.0f}".replace(",", " "), 8, True, navy, PP_ALIGN.CENTER)
        elif t == "table":
            headr = s["head"] or (s["rows"][0] if s["rows"] else ["—"])
            body = s["rows"] if s["head"] else s["rows"][1:]
            ncol = max(len(headr), max((len(r) for r in body), default=1))
            nrow = min(len(body), 12) + 1
            shape = sl.shapes.add_table(nrow, ncol, Inches(0.8), top, SW - Inches(1.6),
                                        Inches(0.4) * nrow)
            tbl = shape.table
            for i in range(ncol):
                cell = tbl.cell(0, i)
                cell.text = str(headr[i]) if i < len(headr) else ""
                for p in cell.text_frame.paragraphs:
                    for r_ in p.runs:
                        r_.font.size, r_.font.bold, r_.font.name = Pt(11), True, "Sylfaen"
            for ri, row in enumerate(body[:nrow - 1], 1):
                for ci in range(ncol):
                    cell = tbl.cell(ri, ci)
                    cell.text = str(row[ci]) if ci < len(row) else ""
                    for p in cell.text_frame.paragraphs:
                        for r_ in p.runs:
                            r_.font.size, r_.font.name = Pt(10), "Sylfaen"
        elif t == "image":
            try:
                p = attachment_path(s["image"]) if s["image"] else None
            except ValueError:
                p = None
            if p:
                from PIL import Image
                iw, ih = Image.open(p).size
                maxw, maxh = SW - Inches(1.6), SH - top - Inches(1.2)
                scale = min(maxw / iw, maxh / ih)
                w, h = Emu(int(iw * scale)), Emu(int(ih * scale))
                sl.shapes.add_picture(str(p), Emu(int((SW - w) / 2)), top, w, h)
                if s["caption"]:
                    box(sl, Inches(0.8), SH - Inches(1.1), SW - Inches(1.6), Inches(0.4), s["caption"],
                        10, False, muted, PP_ALIGN.CENTER)
            else:
                rect(sl, Inches(0.8), top, SW - Inches(1.6), Inches(4), softc)
                box(sl, Inches(0.8), top + Inches(1.8), SW - Inches(1.6), Inches(0.5),
                    "Նկարը ընտրված չէ", 14, False, muted, PP_ALIGN.CENTER)
        elif t == "photo_text":
            p = _img_path(s["image"])
            tw = (SW - Inches(1.6)) * (0.5 if p else 1)
            items = [x for x in s["items"] if x.strip()]
            if items:
                y = top
                for i, it in enumerate(items[:7], 1):
                    rect(sl, Inches(0.8), y, Inches(0.3), Inches(0.3), violet)
                    box(sl, Inches(0.8), y + Inches(0.02), Inches(0.3), Inches(0.3), f"{i:02d}", 8, True,
                        RGBColor(255, 255, 255), PP_ALIGN.CENTER)
                    box(sl, Inches(1.2), y - Inches(0.04), tw - Inches(0.5), Inches(0.6), it, 13)
                    y += Inches(0.62)
            else:
                box(sl, Inches(0.8), top, tw, SH - top - Inches(1), s["body"], 14)
            if p:
                x = Inches(0.8) + tw + Inches(0.35)
                w, hgt = SW - Inches(0.8) - x, SH - top - Inches(1)
                sl.shapes.add_picture(_crop_tmp(p, w / hgt), x, top, w, hgt)
                if s["caption"]:
                    rect(sl, x, top + hgt - Inches(0.35), w, Inches(0.35), navy)
                    box(sl, x + Inches(0.1), top + hgt - Inches(0.35), w - Inches(0.2), Inches(0.3), s["caption"],
                        10, False, RGBColor(255, 255, 255))
        elif t == "gallery":
            ims = [(im, _img_path(im["path"])) for im in s["images"]]
            ims = [(im, p) for im, p in ims if p][:6]
            n = len(ims) or 1
            cols = 1 if n == 1 else 2 if n in (2, 4) else 3
            rows_n = (n + cols - 1) // cols
            gap = Inches(0.2)
            cw = (SW - Inches(1.6) - gap * (cols - 1)) / cols
            chh = (SH - top - Inches(0.9) - gap * (rows_n - 1)) / rows_n
            for i, (im, p) in enumerate(ims):
                r, c = divmod(i, cols)
                x, y = Inches(0.8) + c * (cw + gap), top + r * (chh + gap)
                sl.shapes.add_picture(_crop_tmp(p, cw / chh), Emu(int(x)), Emu(int(y)), Emu(int(cw)), Emu(int(chh)))
                if im.get("caption"):
                    rect(sl, Emu(int(x)), Emu(int(y + chh - Inches(0.32))), Emu(int(cw)), Inches(0.32), navy)
                    box(sl, Emu(int(x + Inches(0.08))), Emu(int(y + chh - Inches(0.33))), Emu(int(cw - Inches(0.16))),
                        Inches(0.3), im["caption"], 9, False, RGBColor(255, 255, 255))
        elif t == "files":
            names = s["items"] or [a["name"] for a in d["attachments"]] or ["—"]
            y = top
            for i, name in enumerate(names[:14], 1):
                rect(sl, Inches(0.8), y, SW - Inches(1.6), Inches(0.38), softc if i % 2 else RGBColor(255, 255, 255))
                box(sl, Inches(1.0), y + Inches(0.04), Inches(0.6), Inches(0.3), f"{i:02d}", 9, True, violet)
                box(sl, Inches(1.7), y + Inches(0.04), SW - Inches(3), Inches(0.3), name, 11, True, ink)
                y += Inches(0.45)
        foot(sl, n)
    prs.save(str(out_path))
    return Path(out_path)
