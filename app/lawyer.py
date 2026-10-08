# -*- coding: utf-8 -*-
"""⚖️ Իրավաբան՝ փաստաթղթերի խիստ համեմատում (նույնիսկ 1 տառ, բացատ կամ կետադրական նշան).

Աջակցվում է՝ .docx, .pdf, .txt, .xlsx
Ինչ է ստուգվում՝
  • ջնջված / ավելացված / փոխված պարբերություններ, աղյուսակների վանդակներ, էջագլուխ/էջատակ
  • փոխված բառեր և ՏԱՌԵՐ (ցույց է տրվում կոնկրետ տառը և դրա Unicode կոդը)
  • բացատներ, անտեսանելի նշաններ (U+00A0, U+200B...), նման տեսք ունեցող տարբեր նշաններ (։ և :)
  • ձևաչափ (թավ, շեղ, ընդգծված, չափ)՝ docx ↔ docx դեպքում
  • ⚠️ կարևոր՝ թվեր, գումարներ, ամսաթվեր, %, ՀՎՀՀ, հաշվեհամար, ժամկետներ, տույժեր
Ոչինչ չի փոխվում ֆայլերում, միայն կարդացվում է."""
import difflib
import html
import re
import unicodedata
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent   # նախագծի արմատը (app/-ից մեկ մակարդակ վեր)
KEYWORDS = ("ՀՎՀՀ", "հաշվեհամար", "դրամ", "տույժ", "տուգանք", "ժամկետ", "վճար", "գումար", "ԱԱՀ", "%",
            "պատասխանատվ", "լուծ", "դադարեց", "երաշխ", "իրավունք", "պարտավոր", "սակագ", "զեղչ")
INVISIBLE = {"\u00a0": "չընդհատվող բացատ", "\u200b": "զրոյական լայնության բացատ", "\u200c": "ZWNJ",
             "\u200d": "ZWJ", "\ufeff": "BOM", "\t": "տաբուլյացիա", "\u2009": "բարակ բացատ", "\u202f": "նեղ բացատ"}


# ------------------------------------------------------------------ reading
class Unit:
    """Համեմատման միավոր՝ պարբերություն/վանդակ/տող + որտեղ է գտնվում."""
    __slots__ = ("text", "where", "fmt")

    def __init__(self, text, where, fmt=None):
        self.text, self.where, self.fmt = text, where, fmt


def _run_fmt(p):
    """Յուրաքանչյուր նշանի ձևաչափը՝ (թավ, շեղ, ընդգծված, չափ)."""
    out = []
    for r in p.runs:
        f = r.font
        sig = (bool(r.bold), bool(r.italic), bool(r.underline), f.size.pt if f.size else None)
        out += [sig] * len(r.text)
    return out


def _docx_units(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = docx.Document(str(path))
    units, pn, tn = [], 0, 0

    def para(p, where):
        units.append(Unit(p.text, where, _run_fmt(p)))

    for el in doc.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            pn += 1
            para(Paragraph(el, doc), f"պարբերություն {pn}")
        elif tag == "tbl":
            tn += 1
            t = Table(el, doc)
            seen = set()
            for ri, row in enumerate(t.rows, 1):
                for ci, cell in enumerate(row.cells, 1):
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    for k, p in enumerate(cell.paragraphs, 1):
                        para(p, f"աղյուսակ {tn}, տող {ri}, սյուն {ci}" + (f", պարբ. {k}" if len(cell.paragraphs) > 1 else ""))
    for si, sec in enumerate(doc.sections, 1):
        for name, part in (("էջագլուխ", sec.header), ("էջատակ", sec.footer)):
            try:
                if part.is_linked_to_previous and si > 1:
                    continue
                for k, p in enumerate(part.paragraphs, 1):
                    para(p, f"{name} (բաժին {si}), տող {k}")
            except Exception:
                pass
    return units


def _pdf_units(path):
    from pypdf import PdfReader
    units = []
    for pi, page in enumerate(PdfReader(str(path)).pages, 1):
        text = page.extract_text() or ""
        for li, line in enumerate(text.splitlines(), 1):
            units.append(Unit(line, f"էջ {pi}, տող {li}"))
    return units


def _xlsx_units(path):
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    units = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    units.append(Unit(str(c.value), f"թերթ «{ws.title}», վանդակ {c.coordinate}"))
    return units


def _txt_units(path):
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1251"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    return [Unit(l, f"տող {i}") for i, l in enumerate(text.splitlines(), 1)]


READERS = {".docx": _docx_units, ".pdf": _pdf_units, ".xlsx": _xlsx_units, ".xlsm": _xlsx_units, ".txt": _txt_units}
SUPPORTED = tuple(READERS)


def read_units(path):
    ext = Path(path).suffix.lower()
    if ext not in READERS:
        raise ValueError("Աջակցվում են՝ .docx, .pdf, .txt, .xlsx (.doc-ը պահպանեք որպես .docx)")
    return READERS[ext](path)


def _reflow(units):
    """Տարբեր ձևաչափեր (pdf ↔ docx)՝ բաժանում ենք նախադասությունների, որ տողադարձերը չխանգարեն."""
    text = " ".join(u.text.strip() for u in units if u.text.strip())
    text = re.sub(r"[ \t]+", " ", text)
    parts = re.split(r"(?<=[։.!?;])\s+", text)
    return [Unit(p, f"նախադասություն {i}") for i, p in enumerate(parts, 1) if p]


# ------------------------------------------------------------------ diff helpers
TOKEN = re.compile(r"\s+|\w+|[^\w\s]", re.U)


def _tokens(s):
    return TOKEN.findall(s)


def _cp(ch):
    try:
        name = unicodedata.name(ch)
    except ValueError:
        name = "?"
    return f"U+{ord(ch):04X} {name}"


def _describe_chars(a, b):
    """Կարճ փոփոխությունների համար՝ կոնկրետ նշանները և կոդերը (անտեսանելի և նման նշանների համար)."""
    notes = []
    for ch in set(a) | set(b):
        if ch in INVISIBLE:
            notes.append(f"«{INVISIBLE[ch]}» ({_cp(ch)})")
    if a and b and len(a) <= 3 and len(b) <= 3:
        same_look = unicodedata.normalize("NFKC", a) == unicodedata.normalize("NFKC", b) or \
            {a, b} in ({"։", ":"}, {"՝", "`"}, {"-", "–"}, {"-", "—"}, {"«", '"'}, {"»", '"'})
        if not same_look and (a + b).isalnum():
            return notes  # սովորական տառ/թիվ՝ կոդերն ավելորդ են
        codes = f"«{a}» [{', '.join(_cp(c) for c in a)}] → «{b}» [{', '.join(_cp(c) for c in b)}]"
        notes.append(("⚠️ ՏԵՍՔՈՎ ՆՈՒՅՆՆ Է, ԲԱՅՑ ՏԱՐԲԵՐ ՆՇԱՆ է՝ " if same_look else "") + codes)
    return notes


def _hidden(text):
    return [f"Պարունակում է անտեսանելի նշան՝ «{INVISIBLE[ch]}» ({_cp(ch)})" for ch in sorted(set(text)) if ch in INVISIBLE]


def _important(*texts):
    t = " ".join(texts)
    return bool(re.search(r"\d", t)) or any(k.lower() in t.lower() for k in KEYWORDS)


def _vis(s):
    """Բացատները տեսանելի դարձնել՝ ցուցադրման համար."""
    if s == "":
        return ""
    if s.strip() == "":
        return "".join("␣" if c == " " else f"[{INVISIBLE.get(c, _cp(c))}]" for c in s)
    return s


def inline_diff(a, b):
    """Բառային + տառային տարբերություն -> [(op, old, new)], op ∈ equal/delete/insert/replace."""
    ta, tb = _tokens(a), _tokens(b)
    out = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, ta, tb, autojunk=False).get_opcodes():
        A, B = "".join(ta[i1:i2]), "".join(tb[j1:j2])
        if op == "replace" and i2 - i1 == 1 and j2 - j1 == 1:  # մեկ բառ՝ տառային մակարդակով
            for op2, x1, x2, y1, y2 in difflib.SequenceMatcher(None, A, B, autojunk=False).get_opcodes():
                out.append((op2, A[x1:x2], B[y1:y2]))
        else:
            out.append((op, A, B))
    return out


def _word_at(text, i, j):
    s = text.rfind(" ", 0, i) + 1
    e = text.find(" ", j)
    return text[s:e if e != -1 else len(text)]


def _details(a, b):
    """Մանրամասն ցուցակ՝ ինչ տառ/բառ է փոխվել, ինչով."""
    ops = inline_diff(a, b)
    lines, pa, pb = [], 0, 0
    for op, x, y in ops:
        if op == "equal":
            pa += len(x)
            pb += len(y)
            continue
        word_old = _word_at(a, pa, pa + len(x)).strip()
        word_new = _word_at(b, pb, pb + len(y)).strip()
        short = len(x) <= 3 and len(y) <= 3 and (x.strip() or y.strip())
        ctx = f" (բառը՝ «{word_old}» → «{word_new}»)" if short and word_old != word_new and (word_old or word_new) else ""
        if op == "delete":
            kind = "տառ" if len(x) == 1 and x.isalpha() else ("բացատ" if not x.strip() else "տեքստ")
            lines.append(f"Ջնջվել է {kind}՝ «{_vis(x)}»{ctx}")
        elif op == "insert":
            kind = "տառ" if len(y) == 1 and y.isalpha() else ("բացատ" if not y.strip() else "տեքստ")
            lines.append(f"Ավելացվել է {kind}՝ «{_vis(y)}»{ctx}")
        else:
            if len(x) == 1 and len(y) == 1 and x.isalpha():
                kind = "տառը"
            elif x.isdigit() and y.isdigit():
                kind = "թվանշանը"
            elif not x.strip() and not y.strip():
                kind = "բացատը"
            elif not x.isalnum() and not y.isalnum() and len(x) <= 3:
                kind = "նշանը"
            else:
                kind = "տեքստը"
            lines.append(f"Փոխվել է {kind}՝ «{_vis(x)}» → «{_vis(y)}»{ctx}")
        lines += ["   " + n for n in _describe_chars(x, y)]
        pa += len(x)
        pb += len(y)
    return lines


def _fmt_changes(ua, ub):
    """Նույն տեքստ, տարբեր ձևաչափ (միայն docx ↔ docx)."""
    if not ua.fmt or not ub.fmt or len(ua.fmt) != len(ub.fmt) or len(ua.fmt) != len(ua.text):
        return []
    names = ("թավ", "շեղ", "ընդգծված", "տառաչափ")
    out, i, n = [], 0, len(ua.fmt)
    while i < n:
        if ua.fmt[i] == ub.fmt[i]:
            i += 1
            continue
        j = i
        while j < n and ua.fmt[j] != ub.fmt[j] and ua.fmt[j] == ua.fmt[i] and ub.fmt[j] == ub.fmt[i]:
            j += 1
        fa, fb = ua.fmt[i], ub.fmt[i]
        what = []
        for k, nm in enumerate(names):
            if fa[k] != fb[k]:
                if k == 3:
                    what.append(f"{nm} {fa[k] or 'ըստ ոճի'} → {fb[k] or 'ըստ ոճի'}")
                else:
                    what.append(f"{nm}՝ {'հանվել է' if fa[k] else 'ավելացվել է'}")
        seg = ua.text[i:j]
        if seg.strip():
            out.append(f"«{seg.strip()[:80]}» — " + ", ".join(what))
        i = j
    return out


# ------------------------------------------------------------------ main compare
def compare(path_a, path_b):
    """-> dict(changes=[...], stats={...}, mode=str). Յուրաքանչյուր change՝
    dict(kind, where_a, where_b, old, new, details[list], important, fmt[list])."""
    ua, ub = read_units(path_a), read_units(path_b)
    ea, eb = Path(path_a).suffix.lower(), Path(path_b).suffix.lower()
    mode = "պարբերություններով"
    if ea != eb:
        ua, ub = _reflow(ua), _reflow(ub)
        mode = "նախադասություններով (ֆայլերը տարբեր ձևաչափի են)"
    A, B = [u.text for u in ua], [u.text for u in ub]
    changes = []
    sm = difflib.SequenceMatcher(None, A, B, autojunk=False)
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            for k in range(i2 - i1):
                f = _fmt_changes(ua[i1 + k], ub[j1 + k])
                if f:
                    changes.append(dict(kind="format", where_a=ua[i1 + k].where, where_b=ub[j1 + k].where,
                                        old=A[i1 + k], new=B[j1 + k], details=[], fmt=f, important=False))
            continue
        if op == "replace":
            # զույգավորում՝ ամենանման պարբերությունները
            left, right = list(range(i1, i2)), list(range(j1, j2))
            pairs, used = [], set()
            for i in left:
                best, bj = 0.0, None
                for j in right:
                    if j in used:
                        continue
                    r = difflib.SequenceMatcher(None, A[i], B[j], autojunk=False).ratio()
                    if r > best:
                        best, bj = r, j
                if bj is not None and best >= 0.45:
                    used.add(bj)
                    pairs.append((i, bj))
            paired_a = {i for i, _ in pairs}
            for i in left:
                if i not in paired_a and A[i].strip():
                    changes.append(dict(kind="delete", where_a=ua[i].where, where_b="", old=A[i], new="",
                                        details=_hidden(A[i]), fmt=[], important=_important(A[i])))
                elif i not in paired_a:
                    changes.append(dict(kind="delete", where_a=ua[i].where, where_b="", old="(դատարկ տող)", new="",
                                        details=[], fmt=[], important=False))
            for i, j in pairs:
                det = _details(A[i], B[j])
                changes.append(dict(kind="change", where_a=ua[i].where, where_b=ub[j].where, old=A[i], new=B[j],
                                    details=det, fmt=_fmt_changes(ua[i], ub[j]) if A[i] == B[j] else [],
                                    important=_important(*[x for op_, x, y in inline_diff(A[i], B[j]) if op_ != "equal"],
                                                         *[y for op_, x, y in inline_diff(A[i], B[j]) if op_ != "equal"])))
            for j in right:
                if j not in used:
                    changes.append(dict(kind="insert", where_a="", where_b=ub[j].where, old="", new=B[j] or "(դատարկ տող)",
                                        details=_hidden(B[j]), fmt=[], important=_important(B[j])))
        elif op == "delete":
            for i in range(i1, i2):
                changes.append(dict(kind="delete", where_a=ua[i].where, where_b="", old=A[i] or "(դատարկ տող)", new="",
                                    details=_hidden(A[i]), fmt=[], important=_important(A[i])))
        elif op == "insert":
            for j in range(j1, j2):
                changes.append(dict(kind="insert", where_a="", where_b=ub[j].where, old="", new=B[j] or "(դատարկ տող)",
                                    details=_hidden(B[j]), fmt=[], important=_important(B[j])))
    ca, cb = sum(len(x) for x in A), sum(len(x) for x in B)
    stats = dict(units_a=len(A), units_b=len(B), chars_a=ca, chars_b=cb,
                 changed=sum(1 for c in changes if c["kind"] == "change"),
                 deleted=sum(1 for c in changes if c["kind"] == "delete"),
                 inserted=sum(1 for c in changes if c["kind"] == "insert"),
                 formatted=sum(1 for c in changes if c["kind"] == "format"),
                 important=sum(1 for c in changes if c["important"]),
                 similarity=round(100 * difflib.SequenceMatcher(None, "\n".join(A), "\n".join(B), autojunk=False).ratio(), 2))
    return dict(changes=changes, stats=stats, mode=mode)


def similarity(path_a, path_b):
    try:
        a = "\n".join(u.text for u in read_units(path_a))
        b = "\n".join(u.text for u in read_units(path_b))
    except Exception:
        return 0.0
    m = difflib.SequenceMatcher(None, a, b, autojunk=False)
    return m.quick_ratio() and m.ratio()


def find_original(path, folders, max_files=500):
    """Ավտոմատ՝ գտնում է ամենանման ֆայլը թղթապանակներում (պատրաստած պայմանագրեր, շաբլոններ).
    Արագ՝ նախ quick_ratio բոլորի համար, ճշգրիտ համեմատությունը՝ միայն լավագույն 8-ի համար."""
    try:
        target = "\n".join(u.text for u in read_units(path))
    except Exception:  # noqa
        return []
    me = Path(path).resolve()
    seen, cands = set(), []
    for f in folders:
        f = Path(f)
        if not f.exists():
            continue
        for p in f.rglob("*"):
            if p.suffix.lower() not in SUPPORTED or p.name.startswith(("~$", "LAW_")):
                continue
            try:
                st = p.stat()
            except OSError:
                continue
            key = (p.name, st.st_size)          # նույն ֆայլի պատճենները (output/ և պահոց) մեկ անգամ
            if key in seen or p.resolve() == me:
                continue
            seen.add(key)
            cands.append((st.st_mtime, p))
    cands.sort(key=lambda x: -x[0])
    rough = []
    for _, c in cands[:max_files]:
        try:
            text = "\n".join(u.text for u in read_units(c))
        except Exception:  # noqa
            continue
        m = difflib.SequenceMatcher(None, target, text, autojunk=False)
        rough.append((m.quick_ratio(), c, m))
    rough.sort(key=lambda x: -x[0])
    best = [(m.ratio(), c) for q, c, m in rough[:8] if q > 0.3]
    best.sort(key=lambda x: -x[0])
    return best[:5]


# ------------------------------------------------------------------ Telegram text (HTML)
def _h(s):
    return html.escape(s, quote=False)


def html_inline(a, b, limit=700):
    parts = []
    for op, x, y in inline_diff(a, b):
        if op == "equal":
            x = x if len(x) < 120 else x[:50] + " … " + x[-50:]
            parts.append(_h(x))
        else:
            if x:
                parts.append(f"<s>{_h(_vis(x))}</s>")
            if y:
                parts.append(f"<b><u>{_h(_vis(y))}</u></b>")
    s = "".join(parts)
    return s if len(s) <= limit else s[:limit] + "…"


ICON = dict(change="✏️ ՓՈԽՎԵԼ Է", delete="❌ ՋՆՋՎԵԼ Է", insert="➕ ԱՎԵԼԱՑՎԵԼ Է", format="🎨 ՁԵՎԱՉԱՓ")


def change_html(n, c):
    where = c["where_a"] or c["where_b"]
    if c["where_a"] and c["where_b"] and c["where_a"] != c["where_b"]:
        where = f"{c['where_a']} → {c['where_b']}"
    head = f"<b>{n}. {ICON[c['kind']]}</b>{'  ⚠️ ԿԱՐԵՎՈՐ' if c['important'] else ''}\n📍 {_h(where)}\n"
    if c["kind"] == "change":
        body = html_inline(c["old"], c["new"]) + "\n" + "\n".join("• " + _h(d) for d in c["details"][:12])
        if len(c["details"]) > 12:
            body += f"\n• … ևս {len(c['details']) - 12} փոփոխություն (տես PDF)"
    elif c["kind"] == "delete":
        body = f"<s>{_h(c['old'][:700])}</s>" + "".join("\n• " + _h(d) for d in c["details"])
    elif c["kind"] == "insert":
        body = f"<b><u>{_h(c['new'][:700])}</u></b>" + "".join("\n• " + _h(d) for d in c["details"])
    else:
        body = _h(c["old"][:200]) + "\n" + "\n".join("• " + _h(f) for f in c["fmt"])
    return head + body


def summary_html(name_a, name_b, res):
    s = res["stats"]
    if not res["changes"]:
        return (f"⚖️ <b>Իրավաբան · արդյունք</b>\n\n📄 Բնօրինակ՝ {_h(name_a)}\n📄 Ստուգվող՝ {_h(name_b)}\n\n"
                f"✅ <b>Տարբերություն ՉԿԱ</b>՝ ոչ մի տառ, բացատ կամ նշան չի փոխվել ({s['chars_a']:,} նշան ստուգված):")
    return (f"⚖️ <b>Իրավաբան · արդյունք</b>\n\n📄 Բնօրինակ՝ {_h(name_a)}\n📄 Ստուգվող՝ {_h(name_b)}\n"
            f"🔍 Համեմատում՝ {res['mode']}\n\n"
            f"Ընդամենը փոփոխություն՝ <b>{len(res['changes'])}</b>\n"
            f"✏️ Փոխված՝ {s['changed']}\n❌ Ջնջված՝ {s['deleted']}\n➕ Ավելացված՝ {s['inserted']}\n"
            f"🎨 Ձևաչափ՝ {s['formatted']}\n⚠️ Կարևոր (թվեր, գումարներ, ժամկետներ...)՝ <b>{s['important']}</b>\n"
            f"📊 Նմանություն՝ {s['similarity']}%\n\n"
            f"Նշանակում՝ <s>ջնջված</s>, <b><u>ավելացված</u></b>. Ամբողջական հաշվետվությունը՝ PDF-ում:")


# ------------------------------------------------------------------ PDF report
def make_report(name_a, name_b, res, out_path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    import mediaplan as mp
    mp._fonts()

    def st(size=9, bold=False, color=colors.black, lead=None):
        return ParagraphStyle("x", fontName="MM-B" if bold else "MM", fontSize=size,
                              leading=lead or size * 1.35, textColor=color)

    def P(t, **k):
        return Paragraph(t, st(**k))

    def inline(a, b):
        out = []
        for op, x, y in inline_diff(a, b):
            if op == "equal":
                out.append(_h(x))
            else:
                if x:
                    out.append(f'<font backColor="#FDE2E2" color="#B42318"><strike>{_h(_vis(x))}</strike></font>')
                if y:
                    out.append(f'<font backColor="#DCFAE6" color="#067647"><u>{_h(_vis(y))}</u></font>')
        return "".join(out)

    s = res["stats"]
    doc = SimpleDocTemplate(str(out_path), pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title="Իրավաբանական համեմատություն")
    story = []
    logo = BASE / "assets" / "logo.png"
    if logo.exists():
        from reportlab.platypus import Image
        from PIL import Image as PI
        w, h = PI.open(logo).size
        img = Image(str(logo), width=34 * mm, height=34 * mm * h / w)
        img.hAlign = "LEFT"
        story += [img, Spacer(1, 4)]
    story += [P("Իրավաբանական համեմատություն", size=16, bold=True, color=colors.HexColor("#1D1B5E")), Spacer(1, 4)]
    info = [["Բնօրինակ", name_a], ["Ստուգվող", name_b], ["Համեմատում", res["mode"]],
            ["Փոփոխություններ", f"{len(res['changes'])} (փոխված {s['changed']}, ջնջված {s['deleted']}, "
                                f"ավելացված {s['inserted']}, ձևաչափ {s['formatted']})"],
            ["Կարևոր", str(s["important"])], ["Նմանություն", f"{s['similarity']}%"],
            ["Նշաններ", f"{s['chars_a']:,} → {s['chars_b']:,}"]]
    t = Table([[P(k, size=8.5, bold=True), P(_h(v), size=8.5)] for k, v in info], colWidths=[34 * mm, 144 * mm])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D0D4E4")),
                           ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#ECE8FD"))]))
    story += [t, Spacer(1, 6),
              P('Նշանակում՝ <font backColor="#FDE2E2" color="#B42318"><strike>ջնջված</strike></font> · '
                '<font backColor="#DCFAE6" color="#067647"><u>ավելացված</u></font> · ␣ = բացատ', size=8), Spacer(1, 8)]
    if not res["changes"]:
        story.append(P("✔ Տարբերություն չկա. ոչ մի տառ, բացատ կամ նշան չի փոխվել։", size=12, bold=True,
                       color=colors.HexColor("#067647")))
    for n, c in enumerate(res["changes"], 1):
        where = c["where_a"] or c["where_b"]
        if c["where_a"] and c["where_b"] and c["where_a"] != c["where_b"]:
            where = f"{c['where_a']} → {c['where_b']}"
        title = {"change": "ՓՈԽՎԵԼ Է", "delete": "ՋՆՋՎԵԼ Է", "insert": "ԱՎԵԼԱՑՎԵԼ Է", "format": "ՁԵՎԱՉԱՓ"}[c["kind"]]
        col = {"change": "#6A3DE8", "delete": "#B42318", "insert": "#067647", "format": "#B4530A"}[c["kind"]]
        rows = [[P(f'<font color="{col}"><b>{n}. {title}</b></font>'
                   f'{"  <font color=\"#B42318\"><b>⚠ ԿԱՐԵՎՈՐ</b></font>" if c["important"] else ""}'
                   f'  <font color="#6B7090">· {_h(where)}</font>', size=9)]]
        if c["kind"] == "change":
            rows += [[P("<b>Էր՝</b> " + _h(c["old"]), size=8.5, color=colors.HexColor("#555A75"))],
                     [P("<b>Դարձել է՝</b> " + _h(c["new"]), size=8.5, color=colors.HexColor("#555A75"))],
                     [P("<b>Տարբերություն՝</b> " + inline(c["old"], c["new"]), size=9)]]
            rows += [[P("• " + _h(d), size=8)] for d in c["details"]]
        elif c["kind"] == "delete":
            rows.append([P(f'<font backColor="#FDE2E2" color="#B42318"><strike>{_h(c["old"])}</strike></font>', size=9)])
            rows += [[P("• " + _h(d), size=8)] for d in c["details"]]
        elif c["kind"] == "insert":
            rows.append([P(f'<font backColor="#DCFAE6" color="#067647"><u>{_h(c["new"])}</u></font>', size=9)])
            rows += [[P("• " + _h(d), size=8)] for d in c["details"]]
        else:
            rows.append([P(_h(c["old"]), size=8.5)])
            rows += [[P("• " + _h(f), size=8)] for f in c["fmt"]]
        card = Table(rows, colWidths=[178 * mm])
        card.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D0D4E4")),
                                  ("LINEBEFORE", (0, 0), (0, -1), 2.5, colors.HexColor(col)),
                                  ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F6F7FB")),
                                  ("LEFTPADDING", (0, 0), (-1, -1), 7), ("TOPPADDING", (0, 0), (-1, -1), 3),
                                  ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
        story += [KeepTogether([card]) if len(rows) < 14 else card, Spacer(1, 5)]
    doc.build(story)
    return Path(out_path)
