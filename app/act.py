# -*- coding: utf-8 -*-
"""ԱԿՏ (Հեռարձակման հաշվետվություն): հաշվարկ ԺԱՄԵՐՈՎ (եթե ֆայլում ժամ կա) և օրերով + լրացում Word շաբլոնից."""
import copy
import csv
import io
import re
from collections import OrderedDict
from datetime import date, datetime, timedelta
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from docfill import MONTHS_GEN, sub_par, to_pdf

BASE = Path(__file__).resolve().parent.parent   # նախագծի արմատը (app/-ից մեկ մակարդակ վեր)
KEYS = dict(
    date=("date", "day", "дата", "день", "ամսաթիվ", "օր", "time", "datetime", "timestamp", "ժամանակ"),
    played=("played", "plays", "count", "воспроизв", "проиграно", "факт", "actual", "հեռարձակ", "նվագարկ", "փաստ"),
    planned=("planned", "plan", "план", "նախատես"),
    clip=("clip", "spot", "ролик", "հոլովակ", "file", "track"),
    status=("status", "статус", "result", "վիճակ"),
    address=("address", "адрес", "հասցե", "shop", "store", "магазин", "խանութ"),
    hour=("time", "время", "час", "hour", "ժամ"),
)
TIME_RE = re.compile(r"(?:^|\s|T)(\d{1,2}):(\d{2})")
BAD = ("fail", "error", "skip", "miss", "not", "no", "err", "не ", "ошиб", "пропущ", "չ", "սխալ", "false", "0")


# ------------------------------------------------------------------ reading
def _norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def _col(headers, kind):
    for i, h in enumerate(headers):
        if any(k in _norm(h) for k in KEYS[kind]):
            return i
    return None


def _to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%y", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M:%S",
              "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return date(int(m[1]), int(m[2]), int(m[3]))
    return None


def _num(v):
    try:
        return int(float(str(v).replace(" ", "").replace(",", ".")))
    except ValueError:
        return None


def read_rows(path):
    path = Path(path)
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        out = []
        for ws in wb:
            out += [list(r) for r in ws.iter_rows(values_only=True)]
        return out
    if ext not in (".csv", ".txt", ".tsv", ""):
        raise ValueError(f"«{path.suffix}» ձևաչափը չի աջակցվում: Ուղարկեք .xlsx կամ .csv")
    raw = path.read_bytes()
    text = None
    for enc in ("utf-8-sig", "utf-16", "cp1251"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:  # ոչ մի կոդավորում չանցավ՝ կարդում ենք կոպիտ, բայց չենք ընկնում
        text = raw.decode("utf-8", "replace")
    if not text.strip():
        raise ValueError("Ֆայլը դատարկ է")
    try:
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t")
    except csv.Error:  # բաժանիչը չճանաչվեց՝ վերցնում ենք ամենահաճախը
        counts = {d: text[:4000].count(d) for d in ",;\t"}
        best = max(counts, key=counts.get)
        dialect = csv.excel
        if counts[best]:
            class _D(csv.excel):
                delimiter = best
            dialect = _D
    return [r for r in csv.reader(io.StringIO(text), dialect)]


def _hour_of(*vals):
    """Բջջի ժամը (datetime/time/«10:20»/«2026-10-01 10:20») -> 0..23 կամ None."""
    for v in vals:
        if isinstance(v, datetime):
            if v.hour or v.minute or v.second:
                return v.hour
            continue
        if hasattr(v, "hour") and hasattr(v, "minute") and not isinstance(v, date):
            return v.hour
        if isinstance(v, float) and 0 < v < 1:          # Excel-ի ժամը՝ օրվա մաս
            return int(round(v * 24 * 60)) // 60 % 24
        m = TIME_RE.search(str(v or ""))
        if m and int(m[1]) < 24 and int(m[2]) < 60:
            return int(m[1])
    return None


def aggregate(rows, start, end, planned_per_day=None, hours_out=None):
    """Վերադարձնում է OrderedDict{date: dict(planned, played, missed)}, clips{name: played}.
    hours_out (dict)՝ լրացվում է ըստ ժամերի {ժամ: [planned, played]}, եթե ֆայլում ժամ կա."""
    hi = next((i for i, r in enumerate(rows) if _col([c for c in r], "date") is not None
               and (_col(r, "played") is not None or len([c for c in r if c]) >= 2)), None)
    if hi is None:
        raise ValueError("Ֆայլում ամսաթվի սյունակ չգտա (date / Ամսաթիվ / Дата)")
    head = rows[hi]
    cd, cp, cl = _col(head, "date"), _col(head, "played"), _col(head, "planned")
    cc, cs = _col(head, "clip"), _col(head, "status")
    ch = next((i for i, h in enumerate(head) if i != cd and any(k in _norm(h) for k in KEYS["hour"])), None)
    days = OrderedDict()
    d = start
    while d <= end:
        days[d] = dict(planned=0, played=0, missed=0)
        d += timedelta(1)
    clips = {}
    for r in rows[hi + 1:]:
        if cd is None or cd >= len(r):
            continue
        dt = _to_date(r[cd])
        if dt is None or dt not in days:
            continue
        rec = days[dt]
        n_played = _num(r[cp]) if cp is not None and cp < len(r) else None
        n_plan = _num(r[cl]) if cl is not None and cl < len(r) else None
        planned_row = 0
        if n_plan is not None:
            rec["planned"] += n_plan
            planned_row = n_plan
        if n_played is not None:
            rec["played"] += n_played
            got = n_played
        else:  # յուրաքանչյուր տող = 1 հեռարձակում (իվենթ)
            bad = cs is not None and cs < len(r) and any(_norm(r[cs]).startswith(b) for b in BAD)
            if bad:
                rec["missed"] += 1
                got = 0
            else:
                rec["played"] += 1
                got = 1
            if n_plan is None:
                rec["planned"] += 1
                planned_row = 1
        if hours_out is not None:
            hr = _hour_of(r[ch] if ch is not None and ch < len(r) else None, r[cd])
            if hr is not None:
                cell = hours_out.setdefault(hr, [0, 0])
                cell[0] += max(planned_row, got)
                cell[1] += got
        if cc is not None and cc < len(r) and r[cc]:
            clips[str(r[cc]).strip()] = clips.get(str(r[cc]).strip(), 0) + got
    for dt, rec in days.items():
        if not rec["planned"] and planned_per_day:
            rec["planned"] = planned_per_day
        if planned_per_day and rec["planned"] < rec["played"] + rec["missed"]:
            rec["planned"] = rec["played"] + rec["missed"]
        rec["missed"] = max(rec["missed"], rec["planned"] - rec["played"])
    return days, clips


def totals(days):
    p = sum(v["planned"] for v in days.values())
    a = sum(v["played"] for v in days.values())
    return p, a, max(p - a, 0)


# ------------------------------------------------------------------ docx
def _set(cell, text):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def _fmt(d):
    return d.strftime("%d.%m.%Y")


def _n(v):
    return f"{v:,}".replace(",", " ")


def make_act(client, contract_no, start, end, days, clips, out_dir, by_addr=None, hours=None):
    doc = docx.Document(BASE / "templates" / "act.docx")
    t0, t1, t2 = doc.tables
    _set(t0.rows[0].cells[1], client)
    _set(t0.rows[1].cells[1], str(contract_no))
    _set(t1.rows[1].cells[0], f"{_fmt(start)} – {_fmt(end)}")
    p, a, miss = totals(days)
    _set(t2.rows[1].cells[1], _n(p))
    _set(t2.rows[1].cells[2], _n(a))
    _set(t2.rows[1].cells[3], _n(miss))
    today = date.today()
    for para in doc.paragraphs:
        if para.text.strip().startswith("Ամսաթիվ"):
            sub_par(para, r"«_+»\s*_+\s*20_+", f"«{today.day:02d}» {MONTHS_GEN[today.month - 1]} {today.year}")

    body = doc.element.body
    anchor = t2._tbl.getnext()  # դատարկ պարբերությունը աղյուսակից հետո
    if anchor is None:  # եթե աղյուսակից հետո ոչինչ չկա՝ ավելացնում ենք խարիսխ պարբերություն
        anchor = doc.add_paragraph()._p

    def table(header, rows, widths_cm):
        t = doc.add_table(rows=1, cols=len(header))
        t.style = t2.style if t2.style else "Table Grid"
        t.autofit = False
        for c, h in zip(t.rows[0].cells, header):
            _set(c, h)
        for r in rows:
            cells = t.add_row().cells
            for c, v in zip(cells, r):
                _set(c, str(v))
        from docx.shared import Cm
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        for ri, row in enumerate(t.rows):
            for ci, c in enumerate(row.cells):
                c.width = Cm(widths_cm[ci])
                for para in c.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci else WD_ALIGN_PARAGRAPH.LEFT
                    for run in para.runs:
                        run.font.size = Pt(9)
                        run.bold = ri == 0
                if ri == 0:
                    tcPr = c._tc.get_or_add_tcPr()
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:fill"), "DDEBF7")
                    tcPr.append(shd)
        borders = OxmlElement("w:tblBorders")
        for b in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{b}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "808080")
            borders.append(e)
        t._tbl.tblPr.append(borders)
        return t

    def caption(text):
        para = doc.add_paragraph()
        r = para.add_run(text)
        r.bold = True
        r.font.size = Pt(10)
        return para

    blocks = []
    if hours:   # հիմնականը՝ ըստ ԺԱՄԵՐԻ
        blocks.append(caption("Աջակցող թվեր՝ ըստ ժամերի"))
        rows = [(h["label"], _n(h["planned"]), _n(h["played"]), _n(h["missed"]),
                 f"{(100 * h['played'] / h['planned']):.0f}%" if h["planned"] else "—") for h in hours]
        hp = sum(h["planned"] for h in hours)
        hm = min(sum(h["missed"] for h in hours), hp)
        ha = hp - hm
        rows.append(("Ընդամենը", _n(hp), _n(ha), _n(hm), f"{(100 * ha / hp):.1f}%" if hp else "—"))
        blocks.append(table(["Ժամ", "Նախատեսված", "Հեռարձակված", "Չհեռարձակված", "Կատարում"], rows,
                            [3.4, 3.2, 3.2, 3.4, 2.6]))
    else:       # ֆայլում ժամ չկա՝ միայն օրերով
        blocks.append(caption("Աջակցող թվեր՝ ըստ օրերի (ֆայլում ժամեր չկան)"))
        rows = [(_fmt(d), _n(v["planned"]), _n(v["played"]), _n(max(v["planned"] - v["played"], 0)),
                 f"{(100 * v['played'] / v['planned']):.0f}%" if v["planned"] else "—") for d, v in days.items()]
        rows.append(("Ընդամենը", _n(p), _n(a), _n(miss), f"{(100 * a / p):.1f}%" if p else "—"))
        blocks.append(table(["Օր", "Նախատեսված", "Հեռարձակված", "Չհեռարձակված", "Կատարում"], rows,
                            [3.4, 3.2, 3.2, 3.4, 2.6]))
    if clips:
        blocks.append(caption("Ըստ հոլովակների (հեռարձակված)"))
        blocks.append(table(["Հոլովակ", "Հեռարձակված"], [(k, _n(v)) for k, v in clips.items()], [10, 5]))
    if by_addr:
        blocks.append(caption("Ըստ հասցեների"))
        arows = [(i, a["addr"], _n(a["planned"]), _n(a["planned"] - a["missed"]), _n(a["missed"]),
                  f"{(100 * (a['planned'] - a['missed']) / a['planned']):.0f}%" if a["planned"] else "—")
                 for i, a in enumerate(by_addr, 1)]
        blocks.append(table(["N", "Հասցե", "Նախատեսված", "Հեռարձակված", "Չհեռարձակված", "Կատարում"], arows,
                            [1.0, 6.4, 2.6, 2.6, 2.6, 1.8]))
        import actpdf
        off = actpdf.off_rows(by_addr)
        blocks.append(caption("Գովազդի անջատումները՝ ըստ հասցեների և օրերի"))
        if off:
            orows = [(r["n"] if r["first"] else "", r["addr"] if r["first"] else "", r["date"], r["hours"],
                      _n(r["missed"])) for r in off]
            blocks.append(table(["N", "Հասցե", "Ամսաթիվ", "Անջատված ժամերը", "Չհեռ."], orows,
                                [1.0, 5.6, 2.4, 6.2, 1.8]))
        else:
            blocks.append(caption("Ընտրված ժամանակահատվածում անջատումներ չեն գրանցվել"))
    for b in blocks:
        el = b._p if hasattr(b, "_p") else b._tbl
        body.remove(el)
        anchor.addnext(el)
        anchor = el
    sp = doc.add_paragraph()
    body.remove(sp._p)
    anchor.addnext(sp._p)
    out_dir = Path(out_dir)
    out_dir.mkdir(exist_ok=True, parents=True)
    safe = re.sub(r"[^\w\-]+", "_", client, flags=re.U)[:30]
    path = out_dir / f"ACT_{safe}_{start:%d%m%Y}_{end:%d%m%Y}.docx"
    doc.save(path)
    return path


def make_act_pdf(*a, **k):  # by_addr՝ kwarg
    path = make_act(*a, **k)
    return path, to_pdf(path)
