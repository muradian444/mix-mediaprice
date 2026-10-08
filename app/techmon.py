# -*- coding: utf-8 -*-
"""📡 Տեխ. մոնիտորինգ՝ ներկառուցված (Google Sheets / Apps Script-ի փոխարեն).

Նախկին «MyMusicMonitoring.gs» սկրիպտի ամբողջ տրամաբանությունը ապրում է հավելվածի ներսում՝
SQLite (data/techmon.db). Drive-ի կարիք չկա:

  բջիջ (օր · օբյեկտ · հասցե · ժամ) -> կարգավիճակ  ⚠ 📞 🟡 🔧 🚧 ✅  (+ նշում, ով/երբ, ավտո/ձեռքով)
  ավտոմատ ⚠ գալիս է նամակներից (techmail.py), մնացածը՝ ձեռքով (մենյուի «Օր» էջ)
  հաշվետվություն ամսվա համար՝ ըստ օբյեկտների, ժամերի, օրերի + համեմատություն նախորդ ամսվա հետ
  ԱԿՏ-ը (monitor.py) տվյալները վերցնում է այստեղից՝ load_period()

Կանոնը նույնն է, ինչ սկրիպտում. ավտոմատ ⚠-ը երբեք չի վերագրում արդեն դրված կարգավիճակը (ձեռքով նշումները պաշտպանված են)."""
import datetime as dt
import io
import re
import sqlite3
import threading
from collections import Counter, OrderedDict, defaultdict
from contextlib import contextmanager

import store
from config import DATA

TZ = dt.timezone(dt.timedelta(hours=4))          # Asia/Yerevan (DST չկա)
DB_PATH = DATA / "techmon.db"
SETTINGS_PATH = DATA / "techmon_settings.json"
LOCK = threading.RLock()

# ------------------------------------------------------------------ կարգավիճակներ
S_OK, S_PROBLEM, S_CONFIRMED, S_CALL, S_SETUP, S_TECH = "✅", "⚠", "🟡", "📞", "🔧", "🚧"
STATUS_LIST = [S_PROBLEM, S_CALL, S_CONFIRMED, S_SETUP, S_TECH, S_OK]
CODE = {S_PROBLEM: "problem", S_CALL: "call", S_CONFIRMED: "confirmed", S_SETUP: "repair", S_TECH: "object"}
MISSED = set(CODE)                                # գովազդը ՉԻ հեռարձակվել (ԱԿՏ-ի հաշվարկով՝ monitor.py-ի հետ նույն)
STATUS_INFO = {                                   # գույները՝ սկրիպտի COLOR_* արժեքները
    S_PROBLEM: dict(bg="#f4cccc", ru="Проблема", hy="Խնդիր", title="❌ СТАТУС: ПРОБЛЕМА (⚠)",
                    text="Проблема, музыка не играет"),
    S_CALL: dict(bg="#f4cccc", ru="Звонок", hy="Զանգ", title="📞 СТАТУС: ЗВОНОК (📞)",
                 text="Менеджер не отвечает/не выходит на связь.\nСтатус мониторинга: отключено"),
    S_CONFIRMED: dict(bg="#fff2cc", ru="Подтверждено", hy="Հաստատված", title="✅ СТАТУС: ПОДТВЕРЖДЕНО (🟡)",
                      text="Подтверждено менеджером, музыка играет"),
    S_SETUP: dict(bg="#cfe2f3", ru="Настройка", hy="Կարգավորում", title="🔧 СТАТУС: НАСТРОЙКА (🔧)",
                  text="Технический отдел совместно с менеджером восстановил проигрывание музыки"),
    S_TECH: dict(bg="#fce5cd", ru="Тех. причина у объекта", hy="Տեխ. պատճառ օբյեկտում",
                 title="🚧 СТАТУС: ТЕХ. ПРИЧИНА У ОБЪЕКТА (🚧)",
                 text="По словам менеджера: света нет / динамики не работают / тех. неполадки"),
    S_OK: dict(bg="#d9ead3", ru="Работает", hy="Աշխատում է", title="✅ СТАТУС: РАБОТАЕТ (✅)",
               text="Музыка работает / проигрывание восстановлено"),
}
MONTHS_RU = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август", "Сентябрь", "Октябрь",
             "Ноябрь", "Декабрь"]
AUTO_MARK = "АВТОМАТИЧЕСКИ"

DEFAULTS = dict(
    enabled=False, imap_host="imap.gmail.com", imap_port=993, imap_user="", imap_password="", imap_folder="INBOX",
    subject="Мониторинг проигрывания музыки", since_days=35, interval_min=10, mark_read=True,
    gmail_label="my_music_processed", work_start=9, work_end=24,
    smtp_host="smtp.gmail.com", smtp_port=465, report_emails="", auto_report=True)


def norm_status(v):
    """Վանդակի արժեքը -> կարգավիճակի նշան (կամ '')."""
    s = str(v or "").replace("️", "").strip()
    if s == "❌":
        s = S_PROBLEM
    return s if s in STATUS_LIST else ""


def clean(s):
    """Օբյեկտի/հասցեի բանալի՝ առանց կրկնակի բացատների (Google-ի աղյուսակներում դրանք դուբլ էին ստեղծում)."""
    return re.sub(r"\s+", " ", str(s or "").replace(" ", " ")).strip()


def now():
    return dt.datetime.now(TZ)


def now_text():
    return now().strftime("%d.%m.%Y %H:%M")


# ------------------------------------------------------------------ կարգավորումներ
def settings(public=False):
    import os
    d = {**DEFAULTS, **(store.read_json(SETTINGS_PATH, {}) or {})}
    if os.getenv("MONITOR_IMAP_PASSWORD"):
        d["imap_password"] = os.getenv("MONITOR_IMAP_PASSWORD")
    for k in ("imap_port", "smtp_port", "since_days", "interval_min", "work_start", "work_end"):
        try:
            d[k] = int(d[k])
        except (TypeError, ValueError):
            d[k] = DEFAULTS[k]
    d["work_start"] = min(max(d["work_start"], 0), 23)
    d["work_end"] = min(max(d["work_end"], d["work_start"]), 24)
    d["interval_min"] = max(d["interval_min"], 1)
    d["since_days"] = min(max(d["since_days"], 1), 400)
    for k in ("enabled", "mark_read", "auto_report"):
        d[k] = bool(d[k])
    if public:
        d["imap_password_set"] = bool(d.pop("imap_password"))
    return d


def save_settings(patch):
    cur = {**DEFAULTS, **(store.read_json(SETTINGS_PATH, {}) or {})}
    for k, v in (patch or {}).items():
        if k not in DEFAULTS:
            continue
        if k == "imap_password" and not str(v or "").strip():
            continue                              # դատարկ դաշտը գաղտնաբառը չի ջնջում
        cur[k] = v.strip() if isinstance(v, str) else v
    store.write_json(SETTINGS_PATH, cur)
    return settings(public=True)


def hours_range():
    s = settings()
    return list(range(s["work_start"], s["work_end"] + 1))


# ------------------------------------------------------------------ տվյալների բազա
_SCHEMA = """
CREATE TABLE IF NOT EXISTS cells(day TEXT NOT NULL, obj TEXT NOT NULL, addr TEXT NOT NULL, hour INTEGER NOT NULL,
  status TEXT NOT NULL, auto INTEGER DEFAULT 0, note TEXT DEFAULT '', user TEXT DEFAULT '', ts TEXT DEFAULT '',
  PRIMARY KEY(day, obj, addr, hour));
CREATE INDEX IF NOT EXISTS cells_day ON cells(day);
CREATE TABLE IF NOT EXISTS history(id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT, day TEXT, obj TEXT, addr TEXT,
  hour INTEGER, old TEXT, new TEXT, user TEXT, source TEXT);
CREATE INDEX IF NOT EXISTS history_cell ON history(day, obj, addr, hour);
CREATE TABLE IF NOT EXISTS objects(obj TEXT NOT NULL, addr TEXT NOT NULL, ext_id TEXT DEFAULT '',
  net_code TEXT DEFAULT '', first_day TEXT DEFAULT '', last_day TEXT DEFAULT '', PRIMARY KEY(obj, addr));
CREATE TABLE IF NOT EXISTS mails(msg_id TEXT PRIMARY KEY, ts TEXT, result TEXT, reason TEXT);
CREATE TABLE IF NOT EXISTS unrec(msg_id TEXT PRIMARY KEY, mail_date TEXT, subject TEXT, reason TEXT, snippet TEXT,
  body TEXT, resolved INTEGER DEFAULT 0);
CREATE TABLE IF NOT EXISTS events(id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT, level TEXT, msg TEXT);
CREATE TABLE IF NOT EXISTS coverage(day TEXT PRIMARY KEY, source TEXT);
CREATE TABLE IF NOT EXISTS kv(k TEXT PRIMARY KEY, v TEXT);
"""
_ready = False


@contextmanager
def db():
    """Կապ՝ ավտոմատ commit/close. Գրառումները սերիալիզացված են (LOCK)."""
    global _ready
    with LOCK:
        con = sqlite3.connect(str(DB_PATH), timeout=30)
        con.row_factory = sqlite3.Row
        try:
            if not _ready:
                con.execute("PRAGMA journal_mode=WAL")
                con.executescript(_SCHEMA)
                have = {r["name"] for r in con.execute("PRAGMA table_info(cells)")}
                for col in ("emoji", "remark", "remark_user"):
                    if col not in have:
                        con.execute(f"ALTER TABLE cells ADD COLUMN {col} TEXT DEFAULT ''")
                _ready = True
            yield con
            con.commit()
        finally:
            con.close()


def kv_get(key, default=""):
    with db() as c:
        r = c.execute("SELECT v FROM kv WHERE k=?", (key,)).fetchone()
    return r["v"] if r else default


def kv_set(key, value):
    with db() as c:
        c.execute("INSERT INTO kv(k,v) VALUES(?,?) ON CONFLICT(k) DO UPDATE SET v=excluded.v", (key, str(value)))


def log(level, msg, con=None):
    """Իրադարձությունների մատյան (նախկին «Central Log»). Պահվում է վերջին 20 000 գրառումը."""
    def _w(c):
        c.execute("INSERT INTO events(ts,level,msg) VALUES(?,?,?)", (now().strftime("%d.%m.%Y %H:%M:%S"), level, str(msg)))
        n = c.execute("SELECT MAX(id) m FROM events").fetchone()["m"] or 0
        if n % 500 == 0:
            c.execute("DELETE FROM events WHERE id <= ?", (n - 20000,))
    if con is not None:
        _w(con)
    else:
        with db() as c:
            _w(c)


def mark_covered(con, day, source="mail"):
    con.execute("INSERT OR IGNORE INTO coverage(day,source) VALUES(?,?)", (str(day), source))


def _touch_object(con, obj, addr, day, ext_id="", net_code=""):
    con.execute("INSERT OR IGNORE INTO objects(obj,addr,ext_id,net_code,first_day,last_day) VALUES(?,?,?,?,?,?)",
                (obj, addr, ext_id or "", net_code or "", str(day), str(day)))
    con.execute("UPDATE objects SET last_day=MAX(last_day,?), first_day=MIN(first_day,?),"
                " ext_id=CASE WHEN ?<>'' THEN ? ELSE ext_id END,"
                " net_code=CASE WHEN ?<>'' THEN ? ELSE net_code END WHERE obj=? AND addr=?",
                (str(day), str(day), ext_id or "", ext_id or "", net_code or "", net_code or "", obj, addr))


def _hour_key(h):
    h = int(h)
    return 24 if h == 0 else h


# ------------------------------------------------------------------ բջիջներ
def manual_note(status, obj, addr, hour, user, auto=False):
    info = STATUS_INFO[status]
    title = info["title"]
    text = info["text"]
    if auto:
        title = f"❌ СТАТУС: ПРОБЛЕМА ({AUTO_MARK}) (⚠)"
        text = "Автоматически обнаружена проблема с воспроизведением музыки"
    return (f"{title}\nОбъект: {obj}\nАдрес: {addr}\nВремя: {_hour_key(hour) % 24:02d}:00\n"
            f"Дата изменения: {now_text()}\nПользователь: {user or 'unknown'}\n---\n{text}")


def get_cell(con, day, obj, addr, hour):
    r = con.execute("SELECT * FROM cells WHERE day=? AND obj=? AND addr=? AND hour=?",
                    (str(day), obj, addr, _hour_key(hour))).fetchone()
    return dict(r) if r else None


def set_cell(day, obj, addr, hour, status, user="", note=None, auto=False, source="manual", ext_id="", net_code="",
             con=None, comment=""):
    """Դնում/մաքրում է բջիջը (status='' -> մաքրել) և գրում պատմություն. -> վանդակի dict կամ None (մաքրված).
    comment՝ օգտատիրոջ ազատ մեկնաբանություն (ավելանում է նշման վերջում)."""
    day = str(day)
    dt.date.fromisoformat(day)
    obj, addr = clean(obj), clean(addr)
    if not obj or not addr:
        raise ValueError("Օբյեկտը և հասցեն պարտադիր են")
    hour = _hour_key(hour)
    if not 0 < hour <= 24:
        raise ValueError("Սխալ ժամ")
    status = norm_status(status)

    def _do(c):
        old = get_cell(c, day, obj, addr, hour)
        old_s = old["status"] if old else ""
        if not status:
            c.execute("DELETE FROM cells WHERE day=? AND obj=? AND addr=? AND hour=?", (day, obj, addr, hour))
            new = None
        else:
            keep_auto = bool(old and old["auto"] and status == S_PROBLEM and AUTO_MARK in (old["note"] or ""))
            n = note if note is not None else (old["note"] if keep_auto else manual_note(status, obj, addr, hour, user, auto))
            if comment and str(comment).strip():
                n += f"\nКомментарий ({user or 'unknown'}): {str(comment).strip()[:500]}"
            c.execute("INSERT INTO cells(day,obj,addr,hour,status,auto,note,user,ts) VALUES(?,?,?,?,?,?,?,?,?) "
                      "ON CONFLICT(day,obj,addr,hour) DO UPDATE SET status=excluded.status, auto=excluded.auto,"
                      " note=excluded.note, user=excluded.user, ts=excluded.ts",
                      (day, obj, addr, hour, status, 1 if (auto or keep_auto) else 0, n,
                       "" if auto else (user or ""), now().isoformat(timespec="seconds")))
            _touch_object(c, obj, addr, day, ext_id, net_code)
            mark_covered(c, day, source)
            new = get_cell(c, day, obj, addr, hour)
        if old_s != status:
            c.execute("INSERT INTO history(ts,day,obj,addr,hour,old,new,user,source) VALUES(?,?,?,?,?,?,?,?,?)",
                      (now().isoformat(timespec="seconds"), day, obj, addr, hour, old_s, status,
                       "" if auto else (user or ""), source))
        return new

    if con is not None:
        return _do(con)
    with db() as c:
        return _do(c)


def set_remark(day, obj, addr, hour, emoji="", text="", user=""):
    """Վանդակի սմայլիկ + մեկնաբանություն (կարգավիճակից անկախ). Երկուսն էլ դատարկ -> հեռացնել."""
    day = str(day)
    dt.date.fromisoformat(day)
    obj, addr, hour = clean(obj), clean(addr), _hour_key(hour)
    emoji, text = str(emoji or "").strip()[:8], str(text or "").strip()[:500]
    with db() as c:
        if not get_cell(c, day, obj, addr, hour):
            raise ValueError("Սկզբում դրեք կարգավիճակ")
        c.execute("UPDATE cells SET emoji=?, remark=?, remark_user=? WHERE day=? AND obj=? AND addr=? AND hour=?",
                  (emoji, text, (user or "") if (emoji or text) else "", day, obj, addr, hour))
        return get_cell(c, day, obj, addr, hour)


def cell_history(day, obj, addr, hour, limit=50):
    with db() as c:
        rows = c.execute("SELECT ts,old,new,user,source FROM history WHERE day=? AND obj=? AND addr=? AND hour=? "
                         "ORDER BY id DESC LIMIT ?", (str(day), clean(obj), clean(addr), _hour_key(hour), limit)).fetchall()
    return [dict(r) for r in rows]


def day_view(day, show_all=False):
    """Օրվա աղյուսակը՝ [{obj, addr, cells:{hour:{s,auto,note,user,ts}}}] (+ ցանկության դեպքում ամբողջ ռեեստրը)."""
    day = str(day)
    dt.date.fromisoformat(day)
    with db() as c:
        cells = c.execute("SELECT obj,addr,hour,status,auto,note,user,ts,emoji,remark,remark_user FROM cells WHERE day=? ORDER BY rowid",
                          (day,)).fetchall()
        known = c.execute("SELECT obj,addr,ext_id,net_code FROM objects ORDER BY obj,addr").fetchall()
        covered = c.execute("SELECT 1 FROM coverage WHERE day=?", (day,)).fetchone() is not None
    rows = OrderedDict()
    for r in cells:
        k = (r["obj"], r["addr"])
        row = rows.setdefault(k, dict(obj=r["obj"], addr=r["addr"], cells={}))
        row["cells"][str(r["hour"])] = dict(s=r["status"], auto=bool(r["auto"]), note=r["note"], user=r["user"], ts=r["ts"],
                                      emoji=r["emoji"] or "", remark=r["remark"] or "", remark_user=r["remark_user"] or "")
    if show_all:
        for r in known:
            rows.setdefault((r["obj"], r["addr"]), dict(obj=r["obj"], addr=r["addr"], cells={}))
    out = list(rows.values())
    for row in out:
        row["problems"] = sum(1 for v in row["cells"].values() if v["s"] == S_PROBLEM)
        row["missed"] = sum(1 for v in row["cells"].values() if v["s"] in MISSED)
    totals = Counter(v["s"] for row in out for v in row["cells"].values())
    return dict(day=day, hours=hours_range(), rows=out, totals=dict(totals), covered=covered or bool(cells),
                known=[dict(obj=r["obj"], addr=r["addr"]) for r in known])


# ------------------------------------------------------------------ ժամանակահատվածներ
def month_bounds(month):
    y, m = (int(x) for x in month.split("-"))
    first = dt.date(y, m, 1)
    last = (first.replace(day=28) + dt.timedelta(4)).replace(day=1) - dt.timedelta(1)
    return first, last


def month_label(month):
    y, m = (int(x) for x in month.split("-"))
    return f"{MONTHS_RU[m - 1]} {y}"


def prev_month(month):
    first, _ = month_bounds(month)
    p = first - dt.timedelta(1)
    return f"{p:%Y-%m}"


def months_with_data():
    """-> [{month, label, days, events}] նորից հինը."""
    with db() as c:
        rows = c.execute("SELECT substr(day,1,7) m, COUNT(DISTINCT day) d, COUNT(*) n FROM cells GROUP BY m").fetchall()
        cov = c.execute("SELECT substr(day,1,7) m, COUNT(*) d FROM coverage GROUP BY m").fetchall()
    data = {r["m"]: dict(month=r["m"], label=month_label(r["m"]), days=r["d"], events=r["n"], covered=0) for r in rows}
    for r in cov:
        data.setdefault(r["m"], dict(month=r["m"], label=month_label(r["m"]), days=0, events=0, covered=0))["covered"] = r["d"]
    return sorted(data.values(), key=lambda x: x["month"], reverse=True)


def covered_days(con, first, last):
    a, b = f"{first:%Y-%m-%d}", f"{last:%Y-%m-%d}"
    days = {r["day"] for r in con.execute("SELECT day FROM coverage WHERE day BETWEEN ? AND ?", (a, b))}
    days |= {r["day"] for r in con.execute("SELECT DISTINCT day FROM cells WHERE day BETWEEN ? AND ?", (a, b))}
    return days


# ------------------------------------------------------------------ ԱԿՏ-ի համար (monitor.py)
def load_period(start, end):
    """Նույն ձևը, ինչ նախկին Google Sheets ընթերցողը. -> (data{date: [{obj, addr, errs{hour: code}}]}, missing[], errors[], files[]).
    Օրը «ծածկված» է, եթե մոնիտորինգն աշխատել է (բջիջ կա կամ ստուգում/ներմուծում է եղել): Ծածկված օր առանց
    խնդրի = ամեն ինչ աշխատել է (դատարկ ցուցակ)."""
    a, b = f"{start:%Y-%m-%d}", f"{end:%Y-%m-%d}"
    with db() as c:
        cells = c.execute("SELECT day,obj,addr,hour,status FROM cells WHERE day BETWEEN ? AND ? ORDER BY day,rowid",
                          (a, b)).fetchall()
        cov = covered_days(c, start, end)
    per_day = defaultdict(OrderedDict)
    for r in cells:
        row = per_day[r["day"]].setdefault((r["obj"], r["addr"]), dict(obj=r["obj"], addr=r["addr"], errs={}))
        if r["status"] in CODE:
            row["errs"][r["hour"]] = CODE[r["status"]]
    data = OrderedDict()
    for day in sorted(cov | set(per_day)):
        data[dt.date.fromisoformat(day)] = list(per_day[day].values()) if day in per_day else []
    months = []
    d = start.replace(day=1)
    while d <= end:
        months.append(f"{d:%Y-%m}")
        d = (d.replace(day=28) + dt.timedelta(4)).replace(day=1)
    have = {x[:7] for x in cov | set(per_day)}
    missing = [m for m in months if m not in have]
    files = [f"Тех мониторинг (встроенный) · {len(data)} дн."] if data else []
    return data, missing, [], files


def months_status(start, end):
    """-> [{month, ready, name, days}] ընտրված ժամանակահատվածի ամիսների համար."""
    have = {m["month"]: m for m in months_with_data()}
    rows = []
    d = start.replace(day=1)
    while d <= end:
        m = f"{d:%Y-%m}"
        h = have.get(m)
        ok = bool(h and (h["events"] or h["covered"]))
        rows.append(dict(month=m, ready=ok, days=(h["covered"] or h["days"]) if h else 0,
                         events=h["events"] if h else 0, name=f"{month_label(m)} · Тех мониторинг" if ok else ""))
        d = (d.replace(day=28) + dt.timedelta(4)).replace(day=1)
    return rows


# ------------------------------------------------------------------ ամսական հաշվետվություն
def report_month(month):
    first, last = month_bounds(month)
    s = settings()
    per_day_hours = s["work_end"] - s["work_start"] + 1
    with db() as c:
        cells = c.execute("SELECT day,obj,addr,hour,status FROM cells WHERE day BETWEEN ? AND ?",
                          (f"{first:%Y-%m-%d}", f"{last:%Y-%m-%d}")).fetchall()
        cov = covered_days(c, first, last)
    today = now().date()
    days_covered = len([d for d in cov if dt.date.fromisoformat(d) <= today])
    by_obj = {}
    hours = defaultdict(Counter)
    by_day = defaultdict(Counter)
    for r in cells:
        st = r["status"]
        o = by_obj.setdefault((r["obj"], r["addr"]), dict(obj=r["obj"], addr=r["addr"], counts=Counter(), days=set(),
                                                           problem_days=set(), hours=Counter()))
        o["counts"][st] += 1
        o["days"].add(r["day"])
        if st in MISSED:
            o["problem_days"].add(r["day"])
        if st == S_PROBLEM:
            o["hours"][r["hour"]] += 1
        hours[r["hour"]][st] += 1
        by_day[r["day"]][st] += 1
    rows = []
    for o in by_obj.values():
        missed = sum(o["counts"][k] for k in MISSED)
        worst = o["hours"].most_common(1)
        rows.append(dict(obj=o["obj"], addr=o["addr"], counts={k: o["counts"].get(k, 0) for k in STATUS_LIST},
                         total=sum(o["counts"].values()), missed=missed, days=len(o["problem_days"]),
                         worst_hour=worst[0][0] if worst else None,
                         uptime=round(100 * (1 - missed / (days_covered * per_day_hours)), 1)
                         if days_covered and per_day_hours else None))
    rows.sort(key=lambda r: (-r["counts"][S_PROBLEM], -r["counts"][S_CALL], -r["missed"], r["obj"]))
    totals = {k: sum(r["counts"][k] for r in rows) for k in STATUS_LIST}
    tops = {}
    for k in STATUS_LIST:
        lst = sorted((r for r in rows if r["counts"][k]), key=lambda r: -r["counts"][k])[:15]
        tops[k] = [dict(obj=r["obj"], addr=r["addr"], n=r["counts"][k], problem=r["counts"][S_PROBLEM], total=r["total"])
                   for r in lst]
    hr = list(range(s["work_start"], s["work_end"] + 1))
    heat = {str(h): {k: hours[h].get(k, 0) for k in STATUS_LIST} for h in hr}
    day_rows = []
    d = first
    while d <= last:
        key = f"{d:%Y-%m-%d}"
        cnt = by_day.get(key, Counter())
        day_rows.append(dict(day=key, label=f"{d:%d.%m}", covered=key in cov,
                             **{k: cnt.get(k, 0) for k in STATUS_LIST}, missed=sum(cnt.get(k, 0) for k in MISSED)))
        d += dt.timedelta(1)
    # համեմատություն նախորդ ամսվա հետ
    pm = prev_month(month)
    pfirst, plast = month_bounds(pm)
    with db() as c:
        pc = c.execute("SELECT status, COUNT(*) n FROM cells WHERE day BETWEEN ? AND ? GROUP BY status",
                       (f"{pfirst:%Y-%m-%d}", f"{plast:%Y-%m-%d}")).fetchall()
    prev = {r["status"]: r["n"] for r in pc}
    compare = []
    for k in STATUS_LIST:
        cur, p = totals[k], prev.get(k, 0)
        compare.append(dict(status=k, cur=cur, prev=p, delta=cur - p,
                            pct=round(100 * (cur - p) / p) if p else (None if not cur else 100)))
    worst_hour = max(heat.items(), key=lambda kv: kv[1][S_PROBLEM], default=(None, {}))
    return dict(month=month, label=month_label(month), prev_month=pm, prev_label=month_label(pm),
                days_in_month=(last - first).days + 1, days_covered=days_covered, hours_per_day=per_day_hours,
                totals=totals, events=sum(totals.values()), missed=sum(totals[k] for k in MISSED),
                objects_affected=len(rows), rows=rows, tops=tops, heat=heat, hours=hr, by_day=day_rows,
                compare=compare, prev_has_data=bool(prev),
                worst_hour=int(worst_hour[0]) if worst_hour[0] and worst_hour[1].get(S_PROBLEM) else None)


# ------------------------------------------------------------------ Excel արտահանում (հին ձևաչափով՝ օրվա թերթեր)
def _fill(hex_):
    from openpyxl.styles import PatternFill
    return PatternFill("solid", fgColor=hex_.lstrip("#").upper())


def export_month_xlsx(month):
    """-> bytes. Թերթեր՝ 📊 ՀԱՇՎԵՏՎՈՒԹՅՈՒՆ, ⏰ ԺԱՄԵՐ, 📈 ՀԱՄԵՄԱՏՈՒԹՅՈՒՆ + ըստ օրերի թերթեր (ՕՕ.ԱԱ.ՏՏՏՏ), ինչպես նախկին ֆայլերում."""
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font
    rep = report_month(month)
    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "📊 ОТЧЁТ ЗА МЕСЯЦ"
    ws["A1"] = f"📊 ОТЧЁТ ЗА МЕСЯЦ — {rep['label']}"
    ws["A1"].font = Font(bold=True, size=14)
    head = ["Объект", "Адрес", *STATUS_LIST, "ВСЕГО", "Дней с проблемой", "Доступность %"]
    ws.append([])
    ws.append(head)
    for c in ws[3]:
        c.font = bold
    for r in rep["rows"]:
        ws.append([r["obj"], r["addr"], *[r["counts"][k] for k in STATUS_LIST], r["total"], r["days"], r["uptime"]])
    ws.append([])
    ws.append(["ИТОГО", "", *[rep["totals"][k] for k in STATUS_LIST], rep["events"]])
    ws.cell(ws.max_row, 1).font = bold
    for k in STATUS_LIST:
        ws.append([])
        ws.append([f"{k} ТОП — {STATUS_INFO[k]['ru']}"])
        ws.cell(ws.max_row, 1).font = bold
        ws.append(["Объект", "Адрес", k, "⚠ (контекст)", "ВСЕГО"])
        for c in ws[ws.max_row]:
            c.font = bold
        for t in rep["tops"][k]:
            ws.append([t["obj"], t["addr"], t["n"], t["problem"], t["total"]])
    ws.freeze_panes = "C4"
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 44

    ws2 = wb.create_sheet("⏰ ОШИБКИ ПО ЧАСАМ")
    ws2["A1"] = f"⏰ ОШИБКИ (⚠) ПО ЧАСАМ — {rep['label']}"
    ws2["A1"].font = Font(bold=True, size=14)
    ws2.append([])
    hrs = rep["hours"]
    ws2.append(["Объект", "Адрес", *[f"{h % 24:02d}:00" for h in hrs], "ИТОГО ⚠"])
    for c in ws2[3]:
        c.font = bold
    ws2.append(["ВСЕГО ПО ЧАСАМ", "", *[rep["heat"][str(h)][S_PROBLEM] for h in hrs], rep["totals"][S_PROBLEM]])
    ws2.cell(4, 1).font = bold
    per_obj = problems_by_hour(month)
    for r in per_obj:
        ws2.append([r["obj"], r["addr"], *[r["hours"].get(h, 0) for h in hrs], r["total"]])
    ws2.freeze_panes = "C5"
    ws2.column_dimensions["A"].width = 44
    ws2.column_dimensions["B"].width = 44

    ws3 = wb.create_sheet("📈 СРАВНЕНИЕ")
    ws3.append(["Показатель", rep["label"], rep["prev_label"], "Δ", "%"])
    for c in ws3[1]:
        c.font = bold
    for r in rep["compare"]:
        ws3.append([f"{r['status']} {STATUS_INFO[r['status']]['ru']}", r["cur"], r["prev"], r["delta"],
                    "—" if r["pct"] is None else f"{r['pct']}%"])
    ws3.column_dimensions["A"].width = 34

    first, last = month_bounds(month)
    with db() as c:
        cells = c.execute("SELECT day,obj,addr,hour,status,note FROM cells WHERE day BETWEEN ? AND ? ORDER BY day,rowid",
                          (f"{first:%Y-%m-%d}", f"{last:%Y-%m-%d}")).fetchall()
    by_day = defaultdict(OrderedDict)
    for r in cells:
        by_day[r["day"]].setdefault((r["obj"], r["addr"]), {})[r["hour"]] = (r["status"], r["note"])
    s = settings()
    for day in sorted(by_day):
        sh = wb.create_sheet(dt.date.fromisoformat(day).strftime("%d.%m.%Y"))
        sh.append(["Объект", "Адрес", *[dt.time(h % 24, 0) for h in range(s["work_start"], s["work_end"] + 1)]])
        for c in sh[1]:
            c.font = bold
            c.alignment = Alignment(horizontal="center")
        for col in range(3, 3 + s["work_end"] - s["work_start"] + 1):
            sh.cell(1, col).number_format = "HH:mm"
            sh.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 7
        sh.column_dimensions["A"].width = 44
        sh.column_dimensions["B"].width = 40
        for (obj, addr), hv in by_day[day].items():
            sh.append([obj, addr])
            row = sh.max_row
            for h, (st, note) in hv.items():
                col = 3 + h - s["work_start"]
                if col < 3:
                    continue
                cell = sh.cell(row, col, st)
                cell.alignment = Alignment(horizontal="center")
                cell.fill = _fill(STATUS_INFO[st]["bg"])
                if note:
                    cell.comment = Comment(note, "Mix Media")
        sh.freeze_panes = "C2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def problems_by_hour(month):
    """[{obj, addr, hours{h: n}, total}] միայն ⚠, նվազման կարգով."""
    first, last = month_bounds(month)
    with db() as c:
        rows = c.execute("SELECT obj,addr,hour,COUNT(*) n FROM cells WHERE status=? AND day BETWEEN ? AND ? "
                         "GROUP BY obj,addr,hour", (S_PROBLEM, f"{first:%Y-%m-%d}", f"{last:%Y-%m-%d}")).fetchall()
    acc = {}
    for r in rows:
        o = acc.setdefault((r["obj"], r["addr"]), dict(obj=r["obj"], addr=r["addr"], hours={}, total=0))
        o["hours"][r["hour"]] = r["n"]
        o["total"] += r["n"]
    return sorted(acc.values(), key=lambda x: -x["total"])


# ------------------------------------------------------------------ ընդհանուր վիճակ
def status():
    s = settings()
    today = now().date()
    with db() as c:
        t = c.execute("SELECT status, COUNT(*) n FROM cells WHERE day=? GROUP BY status", (f"{today:%Y-%m-%d}",)).fetchall()
        objs = c.execute("SELECT COUNT(DISTINCT obj||'|'||addr) n FROM cells WHERE day=?", (f"{today:%Y-%m-%d}",)).fetchone()["n"]
        total_cells = c.execute("SELECT COUNT(*) n FROM cells").fetchone()["n"]
        unrec = c.execute("SELECT COUNT(*) n FROM unrec WHERE resolved=0").fetchone()["n"]
        errs = c.execute("SELECT COUNT(*) n FROM events WHERE level='ERROR' AND id > (SELECT COALESCE(MAX(id),0)-500 FROM events)").fetchone()["n"]
        objects = c.execute("SELECT COUNT(*) n FROM objects").fetchone()["n"]
    return dict(today=f"{today:%Y-%m-%d}", today_totals={r["status"]: r["n"] for r in t}, today_objects=objs,
                cells=total_cells, objects=objects, unrec=unrec, recent_errors=errs,
                configured=bool(s["imap_user"] and s["imap_password"]), enabled=s["enabled"],
                last_run=kv_get("last_run"), last_result=kv_get("last_result"), running=kv_get("running") == "1", progress=kv_get("progress"),
                interval_min=s["interval_min"], months=months_with_data())


def events(level="", limit=200, unrec=False):
    with db() as c:
        if unrec:
            rows = c.execute("SELECT msg_id,mail_date,subject,reason,snippet FROM unrec WHERE resolved=0 "
                             "ORDER BY rowid DESC LIMIT ?", (limit,)).fetchall()
        elif level:
            rows = c.execute("SELECT ts,level,msg FROM events WHERE level=? ORDER BY id DESC LIMIT ?", (level, limit)).fetchall()
        else:
            rows = c.execute("SELECT ts,level,msg FROM events ORDER BY id DESC LIMIT ?", (limit,)).fetchall()
    return [dict(r) for r in rows]
